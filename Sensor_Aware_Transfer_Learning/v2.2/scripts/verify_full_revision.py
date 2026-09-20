"""Independent checks of delivered DOCX tables, pairing, seeds and input hashes."""
import json
import re
import sys
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'code'))
from revision_common import ROOT, SEEDS, sha256, dump
import sensor_drift_transfer as drift

RUN=ROOT/'revision_artifacts/runs/major_revision_001'
OUT=ROOT/'author_review/full_revision_001'
E=OUT/'evidence'

def main():
    checks=[]
    for entry in json.loads((ROOT/'revision_artifacts/environment/input_manifest.json').read_text()):
        assert sha256(Path(entry['path']))==entry['sha256'],entry['path']
    checks.append('All source data hashes, including all ten drift partitions, match the frozen input manifest')
    raw_roles={};raw_configs={};raw_n=0
    for p in (RUN/'jobs').glob('*.json'):
        j=json.loads(p.read_text());i=j['identity'];raw_n+=1
        key=tuple(i[k] for k in ['dataset','fold','seed','budget','design','window'])
        role=(tuple(sorted(j['fit_groups'])),tuple(sorted(j['validation_groups'])))
        if key in raw_roles:assert raw_roles[key]==role,(p.name,'holdout mismatch')
        raw_roles[key]=role
        config=tuple(i[k] for k in ['dataset','fold','budget','design','window','policy'])
        raw_configs.setdefault(config,set()).add(i['seed'])
    assert raw_n==1390 and all(v==set(SEEDS) for v in raw_configs.values())
    checks.append('139 raw-response configurations have ten distinct seeds and common policy holdouts')
    data=drift.load_dataset(drift.DATA);roles={};configs={};drift_n=0
    for p in (RUN/'drift_jobs').glob('*.json'):
        j=json.loads(p.read_text());i=j['metrics'][0];drift_n+=1
        fit=j['fit_indices'];val=j['validation_indices']
        assert all(data['batch'][k]==7 for k in fit+val)
        assert not set(fit)&set(val)
        role=(tuple(sorted(fit)),tuple(sorted(val)));key=(i['seed'],i['budget'])
        if key in roles:assert roles[key]==role,(p.name,'drift holdout mismatch')
        roles[key]=role
        configs.setdefault((i['budget'],i['transform'],i['policy']),set()).add(i['seed'])
        for row in j['predictions']:
            idx=row['exposure']
            assert data['batch'][idx]==row['fold'] and row['fold'] in [8,9,10]
            assert row['gas']==drift.GASES[data['label'][idx]]
            assert np.isclose(row['true'],data['concentration'][idx],rtol=0,atol=0)
    assert drift_n==290 and all(v==set(SEEDS) for v in configs.values())
    checks.append('29 drift configurations have ten seeds; fitting is partition 7 only; test identities/labels and shared holdouts verified')
    doc=Document(OUT/'IEEE_Full_Revision_Clean.docx')
    tables=[[[c.text for c in row.cells] for row in t.rows] for t in doc.tables[1:]]
    manifest=json.loads((OUT/'table_manifest.json').read_text())
    assert len(tables)==len(manifest)==10
    for table,m in zip(tables,manifest):assert table==[m['headers']]+m['rows'],m['title']
    # Independently recompute every primary table value and the full budget grid.
    primary=pd.read_csv(E/'exports/multifold/paired_bootstrap_summary.csv')
    summary=pd.read_csv(E/'exports/10_seed_runs/summary_raw_response.csv')
    def fmt(x,gas):return f'{float(x):.{3 if gas=="Ethylene" else 2}f}'
    for row in tables[0][1:]:
        gas=row[0];a=primary[primary.gas==gas].iloc[0]
        assert row[1:]==[fmt(a.scratch_mae,gas),fmt(a.transfer_mae,gas),fmt(a.improvement,gas),f'[{fmt(a.ci_low,gas)}, {fmt(a.ci_high,gas)}]']
    for row in tables[1][1:]:
        gas,budget=row[:2]
        for cell,policy in zip(row[2:],['scratch','frozen','fine_tuned']):
            q=summary[(summary.gas==gas)&(summary.budget==int(budget))&(summary.policy==policy)&(summary.window=='early')&(summary.design=='farthest')]
            assert len(q)==1 and q.iloc[0].n_seeds==10
            a=q.iloc[0];assert cell==fmt(a.mae_mean,gas)+' ± '+fmt(a.mae_sd,gas)
    # Every number used in other tables is checked against independently read
    # final CSVs or paired ablation differences, not historical snapshots.
    for row in tables[3][1:]:
        policy={'Scratch':'scratch','Extra Trees':'extra_trees','XGBoost':'xgboost','PLSR':'plsr','SVR':'svr','Generic SSL':'generic_ssl','Fine-tuned':'fine_tuned'}[row[0]]
        for cell,gas in zip(row[1:],['NH3','H2','CO','Ethylene']):
            q=summary[(summary.gas==gas)&(summary.budget==(57 if gas in ['NH3','H2'] else 24))&(summary.policy==policy)&(summary.window=='early')&(summary.design=='farthest')]
            a=q.iloc[0];assert len(q)==1 and a.n_seeds==10;assert cell==fmt(a.mae_mean,gas)+' ± '+fmt(a.mae_sd,gas)
    design=pd.read_csv(E/'exports/calibration_design/factorial_summary.csv')
    for row in tables[2][1:]:
        gas,des=row[:2]
        for cell,policy in zip(row[2:],['scratch','fine_tuned']):
            a=design[(design.gas==gas)&(design.design==des.lower())&(design.policy==policy)&(design.budget==5)].iloc[0]
            assert a.n_seeds==10 and cell==fmt(a.mae_mean,gas)+' ± '+fmt(a.mae_sd,gas)
    raw=pd.read_csv(E/'exports/policy_selection/all_policy_results.csv')
    raw=raw[(raw.fold==0)&(raw.design=='farthest')&(raw.window=='early')]
    for row,variant in zip(tables[4][1:],['no_sensor_ids','no_attention','no_temporal_mask','no_sensor_mask']):
        for cell,gas in zip(row[1:],['NH3','H2','CO','Ethylene']):
            d=raw[(raw.gas==gas)&(raw.budget==(57 if gas in ['NH3','H2'] else 24))]
            delta=d[d.policy==variant].set_index('seed').mae-d[d.policy=='fine_tuned'].set_index('seed').mae
            assert len(delta)==10 and cell==fmt(delta.mean(),gas)+' ± '+fmt(delta.std(ddof=1),gas)
    window=pd.read_csv(E/'exports/window_sensitivity/window_position_results.csv')
    for row in tables[5][1:]:
        for cell,(gas,policy) in zip(row[1:],[('NH3','scratch'),('NH3','fine_tuned'),('H2','scratch'),('H2','fine_tuned')]):
            a=window[(window.window==row[0].lower())&(window.gas==gas)&(window.policy==policy)&(window.budget==57)].iloc[0]
            assert a.n_seeds==10 and cell==fmt(a.mae_mean,gas)+' ± '+fmt(a.mae_sd,gas)
    macro=pd.read_csv(E/'exports/sensor_drift/macro_error.csv')
    for row in tables[6][1:]:
        policy={'Constant':'constant','Scratch':'scratch','Frozen':'frozen','Fine-tuned':'fine_tuned'}[row[0]]
        for cell,partition in zip(row[1:],[8,9,10]):
            a=macro[(macro.fold==partition)&(macro.budget==20)&(macro['transform']=='log_huber')&(macro.policy==policy)].mae
            assert len(a)==10 and cell==f'{a.mean():.2f} ± {a.std(ddof=1):.2f}'
    pergas=pd.read_csv(E/'exports/sensor_drift/per_gas_error.csv')
    for row in tables[7][1:]:
        for cell,partition in zip(row[1:],[8,9,10]):
            if row[0]=='Macro':a=macro[(macro.fold==partition)&(macro.budget==20)&(macro['transform']=='log_huber')&(macro.policy=='fine_tuned')].nmae
            else:a=pergas[(pergas.fold==partition)&(pergas.budget==20)&(pergas['transform']=='log_huber')&(pergas.policy=='fine_tuned')&(pergas.gas==row[0])].nmae
            assert len(a)==10 and cell==f'{100*a.mean():.2f} ± {100*a.std(ddof=1):.2f}'
    for row,transform in zip(tables[8][1:],['log_huber','log_fixed_huber','linear_huber','linear_mae']):
        for cell,partition in zip(row[1:],[8,9,10]):
            a=macro[(macro.fold==partition)&(macro.budget==20)&(macro['transform']==transform)&(macro.policy=='fine_tuned')].mae
            assert len(a)==10 and cell==f'{a.mean():.2f} ± {a.std(ddof=1):.2f}'
    for row in tables[9][1:]:
        gas,label=row[:2];policy={'Scratch':'scratch','Frozen':'frozen','Fine-tuned':'fine_tuned'}[label]
        a=summary[(summary.gas==gas)&(summary.budget==(57 if gas in ['NH3','H2'] else 24))&(summary.policy==policy)&(summary.window=='early')&(summary.design=='farthest')].iloc[0]
        assert row[2:]==[f'{a.r2_mean:.3f}',fmt(a.rmse_mean,gas),fmt(a.q90_mean,gas)]
    checks.append('All ten delivered DOCX tables independently match final CSVs or recomputed paired ablation differences')
    response=Document(OUT/'RESPONSE_TO_REVIEWERS.docx')
    rt=[[[c.text for c in row.cells] for row in t.rows] for t in response.tables]
    assert rt==tables
    headings=[p.text for p in response.paragraphs if p.style.name=='Heading 1']
    assert sum(x.startswith('Major comment') for x in headings)==9
    assert sum(x.startswith('Moderate comment') for x in headings)==11
    assert sum(x.startswith('Presentation comment') for x in headings)==6
    checks.append('Response covers nine major, eleven moderate and six presentation points; its numerical tables match the manuscript')
    text='\n'.join(p.text for p in doc.paragraphs)
    body,refs=text.split(' References',1)
    cited=set(re.findall(r'\[(\d+)\]',body))
    for lo,hi in re.findall(r'\[(\d+)\][–-]\[(\d+)\]',body):cited.update(map(str,range(int(lo),int(hi)+1)))
    listed=set(re.findall(r'^\[(\d+)\]',refs,re.M))
    assert cited==listed==set(map(str,range(1,31))),(cited,listed)
    checks.append('All 30 references are cited; no undefined citations')
    validation=json.loads((OUT/'document_validation.json').read_text())
    assert validation['original_equation_objects_preserved']==64
    dump(OUT/'full_revision_validation.json',dict(status='PASS',checks=checks,raw_jobs=raw_n,drift_jobs=drift_n,
        source_docx_sha256=validation['original_sha256'],private_characterization='Harvard Dataverse dataset citation, design reference and companion Nano MEMS description added; dataset-specific settings remain unverified',
        precision_note='Recorded neural/drift Q90 uses NumPy float32 arithmetic. The consistency checker reproduces this arithmetic; scientific results were not changed.'))
    print('\n'.join('PASS: '+c for c in checks))

if __name__=='__main__':main()
