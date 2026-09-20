"""Build result-grounded review draft, figures, claims, and response after reruns."""
import argparse
import json
import re
import sys
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'code'))
from revision_common import ROOT, dump

ADDITIONS = [
    'Y. Nie, N. H. Nguyen, P. Sinthong, and J. Kalagnanam, "A Time Series is Worth 64 Words: Long-term Forecasting with Transformers," ICLR, 2023. https://arxiv.org/abs/2211.14730',
    'G. Zerveas, S. Jayaraman, D. Patel, A. Bhamidipaty, and C. Eickhoff, "A Transformer-based Framework for Multivariate Time Series Representation Learning," KDD, 2021. https://arxiv.org/abs/2010.02803',
    'Z. Yue et al., "TS2Vec: Towards Universal Representation of Time Series," AAAI, 2022. https://arxiv.org/abs/2106.10466',
    'S. Wold, M. Sjostrom, and L. Eriksson, "PLS-regression: a basic tool of chemometrics," Chemometrics and Intelligent Laboratory Systems, vol. 58, pp. 109–130, 2001. https://doi.org/10.1016/S0169-7439(01)00155-1',
    'A. J. Smola and B. Scholkopf, "A tutorial on support vector regression," Statistics and Computing, vol. 14, pp. 199–222, 2004. https://doi.org/10.1023/B:STCO.0000035301.49549.88',
    'R. W. Kennard and L. A. Stone, "Computer Aided Design of Experiments," Technometrics, vol. 11, pp. 137–148, 1969. https://doi.org/10.1080/00401706.1969.10490666',
    'R. F. Barber, E. J. Candes, A. Ramdas, and R. J. Tibshirani, "Conformal prediction beyond exchangeability," Annals of Statistics, vol. 51, 2023. https://arxiv.org/abs/2202.13415',
    'R. J. Tibshirani, R. F. Barber, E. J. Candes, and A. Ramdas, "Conformal Prediction Under Covariate Shift," NeurIPS, 2019. https://arxiv.org/abs/1904.06019',
    'P. Geurts, D. Ernst, and L. Wehenkel, "Extremely randomized trees," Machine Learning, vol. 63, pp. 3–42, 2006. https://doi.org/10.1007/s10994-006-6226-1',
    'T. Chen and C. Guestrin, "XGBoost: A Scalable Tree Boosting System," KDD, pp. 785–794, 2016. https://arxiv.org/abs/1603.02754',
]


class Writer:
    def __init__(self, run):
        self.run=run;self.claims=[];self.parts=[]

    def add(self,text):self.parts.append(text)

    def table(self,relative,frame,keys,columns,title):
        if frame.empty:
            self.add(f'{title}: results not yet available.');return
        headers=[label for col,label in columns]
        lines=[title,'','| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']
        for _,row in frame.iterrows():
            cells=[];pending=[]
            where={key:str(row[key]) for key in keys}
            for col,label in columns:
                value=row[col]
                if isinstance(value,(float,np.floating)) and pd.notna(value):
                    rendered=f'{value:.2f}'
                    pending.append(dict(csv=relative,where=where,column=col,decimals=2,rendered=rendered))
                elif isinstance(value,(int,np.integer)):
                    rendered=str(int(value))
                    pending.append(dict(csv=relative,where=where,column=col,decimals=0,rendered=rendered))
                else:rendered='undefined' if pd.isna(value) else str(value)
                cells.append(rendered)
            line='| '+' | '.join(cells)+' |';lines.append(line)
            self.claims.extend([c|{'text':line} for c in pending])
        self.add('\n'.join(lines))

    def sentence(self,relative,row,where,column,label,decimals=2):
        value=f'{float(row[column]):.{decimals}f}'
        text=f'{label} {value}'
        self.claims.append(dict(csv=relative,where=where,column=column,decimals=decimals,rendered=value,text=text))
        return text


def figures(run,out):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.size':11,'axes.labelsize':11,'legend.fontsize':9,'pdf.fonttype':42,'svg.fonttype':'none'})
    fig,axes=plt.subplots(1,3,figsize=(7.2,2.7))
    for ax,dataset,pattern in zip(axes,['NH3/H2','CO/ethylene','Sensor_Drift'],
                   ['nh3_h2_f0_early_s*_full.json','co_ethylene_f0_early_s*_full.json','drift_s*.json']):
        curves=[json.loads(p.read_text())['loss'] for p in (run/'checkpoints').glob(pattern)]
        if curves:
            values=np.asarray(curves);mean=values.mean(0);sd=values.std(0,ddof=1) if len(values)>1 else mean*0
            epoch=np.arange(1,len(mean)+1)
            ax.plot(epoch,mean,label=f'{len(values)} seeds');ax.fill_between(epoch,mean-sd,mean+sd,alpha=.2)
            ax.legend();ax.set(title=dataset,xlabel='Epoch')
    axes[0].set_ylabel('Reconstruction loss');fig.tight_layout()
    for extension in ['pdf','svg','png']:fig.savefig(out/f'fig3_pretraining.{extension}',dpi=200)
    plt.close(fig)
    # Preserve the supplied Fig. 2 image verbatim from its original paragraph.
    source=next((ROOT/'revision_artifacts/environment').glob('*.docx'))
    original=Document(source)
    paragraph=original.paragraphs[86]
    for blip in paragraph._p.xpath('.//a:blip'):
        rid=blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
        part=original.part.related_parts[rid]
        suffix=Path(str(part.partname)).suffix
        (out/f'fig2_preserved{suffix}').write_bytes(part.blob)


def build(run,allow_incomplete=False):
    required={'fixed','multifold','classical','generic','design','window','ablation','drift'}
    completed={p.stem.removesuffix('_completed') for p in run.glob('*_completed.json')}
    protocol=json.loads((run/'protocol.json').read_text())
    complete=required<=completed and not protocol['smoke']
    if not complete and not allow_incomplete:
        raise RuntimeError('Full reruns are incomplete; final abstract must wait. Use --allow-incomplete only for an explicitly interim draft.')
    out=run/'manuscript';out.mkdir(parents=True,exist_ok=True)
    w=Writer(run)
    def read(relative):
        path=run/'exports'/relative
        if not path.exists() or path.stat().st_size<5:return pd.DataFrame()
        return pd.read_csv(path)
    primary=read('multifold/paired_bootstrap_summary.csv')
    w.add('# Sensor-Aware Self-Supervised Transfer Learning for Label-Efficient Multigas Electronic Nose Calibration')
    if not complete:w.add('**INTERIM REVIEW DRAFT: full experiment matrix is incomplete; no final abstract or publication claim is certified.**')
    w.add('## Abstract')
    if complete:
        sentences=[]
        for _,r in primary.iterrows():
            where={k:str(r[k]) for k in ['dataset','budget','design','window','gas']}
            sentences.append(w.sentence('multifold/paired_bootstrap_summary.csv',r,where,'relative_improvement_pct',f"{r.gas}: relative MAE change favoring fine-tuning =")+'%.')
        w.add('We study reduced labeled-calibration requirements for same-system multigas electronic-nose recalibration using masked reconstruction of sensor responses. '
              'Pretraining withholds existing exposure labels and does not demonstrate reduced gas-generation cost. A shared temporal encoder, sensor-identity embeddings, '
              'and cross-sensor attention are evaluated using a predefined fine-tuned transfer policy and 10 seeds. Primary evidence comes from six held-partition NH3/H2 '
              'folds and three chronological CO/ethylene folds. '+' '.join(sentences)+' Paired uncertainty is reported in Table 1. '
              'Fixed-split budgets, random calibration design, classical regression, a PatchTST-style generic SSL control, and window sensitivity qualify the results. '
              'Sensor_Drift concentration regression is oracle-class conditional; partition 10 is interpreted using paired uncertainty. '
              'The experiments do not establish cross-device transfer or uniformly monotonic budget effects.')
    else:w.add('The final abstract will be generated only after all full-length reruns complete.')
    w.add('Keywords: electronic nose; concentration regression; self-supervised learning; sensor drift; calibration design; transfer learning.')
    w.add('## I. Introduction\n\nControlled multigas exposures provide quantitative calibration data, but collecting them remains an experimental burden. '
          'This study asks whether a previously acquired pool of responses can reduce the number of concentration labels used during later same-system calibration. '
          'The pretraining exposures were originally controlled and labeled; withholding their labels is a representation-learning experiment, not evidence that those exposures can be generated more cheaply. '
          'The contribution is evaluated through independent adaptation-label budgets and held-out exposure groups. We distinguish representation transfer, calibration-pair design, and chronological robustness.')
    w.add('## II. Related Work\n\nElectronic-nose semi-supervised recognition includes M-Training [1] and M-S4VMs [2]. Drift work spans domain adaptation, '
          'few-shot learning, graph-based methods, and meta-learning [3] [4] [5] [6] [7]. Quantitative gas estimation has used dynamic models and temporal feature fusion [8] [9] [10], '
          'as well as small-data and cross-gas transfer [11] [12]. Time-series SSL covers several pretext-task families [13]. '
          'PatchTST uses shared channel-independent patch Transformers and masked pretraining [18]; TST uses a Transformer representation-learning framework [19], '
          'whereas TS2Vec learns hierarchical contrastive representations [20]. We implement a compact PatchTST-style control rather than treating internal sensor-ID ablation as the generic SSL baseline.\n\n'
          'Classical controls include partial least-squares regression [21], support-vector regression [22], Extra Trees [26], and XGBoost [27]. '
          'Minimal-experiment sensor-array calibration [14] and space-covering experimental design [23] motivate the farthest-point control. '
          'Conformal prediction under nonexchangeability [24] and covariate shift [25] explain why ordinary split-conformal guarantees cannot be transferred directly to the present chronological deployment setting.')
    w.add('## III. Method\n\nFor a 64-by-S response window X, the same dilated temporal CNN processes each sensor signal and observation mask. '
          'Three convolutional layers have widths 24, 32, and 32, kernels 5, 5, and 3, and dilations 1, 2, and 4. Mean and maximum temporal pooling '
          'are concatenated and projected to a 32-dimensional token. A learned physical channel embedding is added before two four-head cross-sensor Transformer layers '
          'with 96-unit feed-forward sublayers. The supervised head maps flattened sensor tokens through widths 128 and 32 to two concentration outputs.\n\n'
          'The pretext loss is masked-entry MSE plus 0.05 times visible-entry MSE. Masking removes temporal spans and entire sensor channels. '
          'AdamW uses 60 pretraining epochs for raw responses. Supervised adaptation uses 120 epochs and Huber loss on fixed-scale targets, with delta 0.1. '
          'Scratch starts with a random encoder; frozen transfer updates only the head; fine-tuning uses encoder and head learning rates of 0.00008 and 0.0008. '
          'Fine-tuning is predefined as the primary transfer policy for every gas. Each seed selects a common exposure-group adaptation holdout; only that holdout selects checkpoints.\n\n'
          'Farthest-point calibration design sequentially maximizes minimum distance in normalized concentration coordinates. The initial pair is seeded randomly, '
          'and both farthest and random designs select unique concentration pairs. These coordinates are assumed known during experimental design. '
          'The procedure is not label acquisition from an unknown concentration pool. Random design uses the same budgets and seed list. '
          'Both designs are evaluated with scratch and fine-tuned training to separate design and transfer effects.')
    w.add('## IV. Data and Evaluation Protocol\n\nThe private NH3/H2 workbook contains 20 response channels, 378 exposure identifiers, and six consecutive partitions inferred '
          'from groups of 63 identifiers. These partitions must not be called verified acquisition days. The primary protocol holds out each partition in turn, '
          'uses its preceding cyclic partition for adaptation, and uses the remaining four for label-free pretraining. This evaluates held-partition generalization; '
          'some pretraining partitions occur after the held-out partition, so the six-fold analysis is not a prospective forecasting experiment. '
          'The secondary fixed split uses partitions 1–4, 5, and 6 for pretraining, adaptation, and test. Budgets are 5, 10, 20, 40, and 57 pairs.\n\n'
          'The public dynamic gas mixture dataset [15] supplies 16 channels and 245 cached exposure groups. The primary protocol divides acquisition order into seven blocks '
          'and uses three rolling pretrain–adapt–test folds, with three pretraining blocks followed by one adaptation and one test block. Full coverage is the number of '
          'available unique adaptation pairs in each fold. The secondary five-block protocol uses the first three blocks, fourth block, and fifth block respectively; '
          'budgets are 5, 10, 20, and 24 pairs.\n\n'
          'Pretraining robust scalers use response values only. All five offset windows from each fitting exposure remain in the same role, and tests use one window per exposure. '
          'Every adaptation budget includes both fitting and held-out labels. An audit of 623 identified exposures/windows found no identical signal fingerprints and no exposure-ID '
          'overlap in 11 split definitions. Distinct identifiers do not independently verify physical acquisition independence; laboratory logs remain necessary. '
          'The CO/ethylene fingerprint audit covers cached windows, not complete raw traces.\n\n'
          'The private-array manufacturer, sensing material, heater settings, physical baseline convention, sampling rate, chamber/flow details, ambient temperature, '
          'relative humidity and its control, and concentration-reference accuracy were unavailable in the inspected acquisition records. This limits reproducibility and '
          'precludes attributing sensor-ID effects to sensing chemistry. Sample count is not converted into elapsed seconds. The 0–500 ppm calibration setting is not '
          'validation of low-concentration ambient environmental monitoring.\n\n'
          'The Sensor_Drift concentration dataset [16] and drift-compensation literature [17] motivate the descriptor extension. Pretraining uses partitions 1–6, adaptation uses '
          'partition 7, and tests are partitions 8–10. Six concentration heads are selected by the true gas identity. This oracle-class conditional problem does not include gas-classification error. '
          'The constant control predicts the per-gas fitting-subset mean. Budgets are 1, 2, 5, 10, and 20 labels/gas. One label/gas uses a fixed final checkpoint because '
          'no disjoint holdout is possible. Smaller budgets use log1p targets divided by log1p(1000), a predefined unit scale. At 20 labels/gas, the historical per-gas '
          'fitting-subset log standardization is compared with fixed-log scaling, linear targets divided by 1000 with Huber loss, and the same linear scale with MAE loss.\n\n'
          'All principal experiments use 10 seeds: 42, 142, 242, 342, 442, 542, 642, 742, 842, and 942. Pretraining is rerun per seed. '
          'We report MAE, RMSE, R2, and Q0.9 of absolute errors. Per-gas drift NMAE is MAE divided by the historical concentration range in partitions 1–6; these reference labels '
          'are used only for reporting, not pretraining, fitting or selection. Some test-partition gas ranges are zero, so a fixed historical reference permits normalization of all six gases. MacroMAE and macro-NMAE weight gases equally. '
          'Paired 95% bootstrap intervals resample seed and fold indices independently with equal fold weights, using 5000 replicates. They are conditional on observed test '
          'exposures; shared training partitions and the small number of folds limit population inference. We do not pool repeated predictions as independent observations.')
    w.add('## V. Results\n\n### V.A. Self-Supervised Pretraining\n\nReconstruction curves are supporting optimization diagnostics, not evidence of concentration accuracy. '
          'Figure 3 shows mean loss with one seed standard deviation; the architecture and split schematic from the supplied Figure 2 is preserved.\n\n'
          'Figure 3. Masked-reconstruction pretraining with 10 seeds; 60 epochs for raw responses and 40 for Sensor_Drift. Shading denotes seed SD. Vector PDF/SVG use print-size labels.')
    w.add('### V.B. Primary Multi-Fold Validation\n\nTable 1 is the primary generalization evidence. Positive differences favor the predefined fine-tuned policy. '
          'Across-fold SD describes variation of fold means, and paired-prediction counts include repeated seeds rather than claiming that all repetitions are independent exposures.')
    pkeys=['dataset','budget','design','window','gas']
    w.table('multifold/paired_bootstrap_summary.csv',primary,pkeys,
            [('gas','Gas'),('scratch_mae','Scratch MAE'),('transfer_mae','Fine-tuned MAE'),('improvement','Paired difference'),
             ('relative_improvement_pct','Relative difference (%)'),('ci_low','95% CI lower'),('ci_high','95% CI upper'),
             ('across_fold_sd','Across-fold SD'),('n_paired_predictions','Paired predictions')],'Table 1. Primary multi-fold results (ppm).')
    w.add('### V.C. Fixed-Split Label-Budget Analysis\n\nTable 2 reports all policies without selecting among them using test errors. '
          'Means and SD quantify variation across 10 seeds. Label-budget effects need not be monotonic for each gas, and the fixed split is secondary to Table 1.')
    summary=read('10_seed_runs/summary_raw_response.csv')
    skeys=['dataset','fold','budget','design','window','policy','gas']
    fixed=summary[(summary.policy.isin(['scratch','frozen','fine_tuned']))&(summary.design=='farthest')&(summary.window=='early')] if len(summary) else summary
    cols=[('gas','Gas'),('budget','Pairs'),('policy','Policy'),('mae_mean','MAE mean'),('mae_sd','MAE SD')]
    w.table('10_seed_runs/summary_raw_response.csv',fixed,skeys,cols,'Table 2. Fixed-split calibration budgets (ppm).')
    w.add('### V.D. Calibration-Design Control\n\nBoth random and farthest-point designs are crossed with scratch and fine-tuned training. '
          'The factorial summary and paired design statistics are supplied with all budgets in the reproducibility artifacts. Table 3 shows the low-budget control; '
          'the design intervention requires known concentration coordinates and does not establish active selection of unlabeled ambient observations.')
    design=read('calibration_design/factorial_summary.csv')
    w.table('calibration_design/factorial_summary.csv',design[design.budget==5] if len(design) else design,skeys,
            [('gas','Gas'),('design','Pair design'),('policy','Training'),('mae_mean','MAE mean'),('mae_sd','MAE SD')],'Table 3. Five-pair calibration design (ppm).')
    w.add('### V.E. Baseline Comparison\n\nTable 4 reports full-budget results; the supplement includes every low and medium budget. '
          'PLSR, RBF SVR, Extra Trees, and XGBoost receive flattened scaled 64-sample windows (1280 or 1024 features). PLSR and SVR additionally standardize fitting features '
          'and tune hyperparameters using adaptation-only validation. The revised tree comparison is rerun with this explicit common representation, rather than reusing historical engineered-feature scores. '
          'The generic control is a compact PatchTST-style reimplementation with eight length-8 patches, a shared channel-independent encoder, temporal positions, and no physical sensor embeddings. '
          'Three patches per channel are masked. A matched supervised head retains channel ordering. It is not an exact reproduction of published forecasting settings; '
          'instance normalization is omitted to match preprocessing and preserve absolute response information. Encoder capacity differs and is reported in the model-size artifact.')
    baselines=read('baselines/baseline_summary.csv')
    if len(baselines):baselines=baselines[(baselines.budget==baselines.groupby('dataset').budget.transform('max'))&baselines.policy.isin(['scratch','extra_trees','xgboost','plsr','svr','generic_ssl','fine_tuned'])]
    w.table('baselines/baseline_summary.csv',baselines,skeys,[('gas','Gas'),('policy','Model'),('mae_mean','MAE mean'),('mae_sd','MAE SD')],'Table 4. Full-budget baseline comparison (ppm).')
    w.add('### V.F. Ablation and Window Sensitivity\n\nTable 5 gives full-budget architectural ablations with seed uncertainty. '
          'Gas-dependent sensor-identity effects may reflect selectivity, sampling variability, or both; the current design cannot separate these explanations. '
          'Early, middle, late, and full-exposure-resampled 64-sample definitions are compared in Table 6 under fixed exposure roles and full calibration budgets. '
          'The early window remains the predefined primary protocol; the sensitivity study does not retroactively select a test-optimal window.')
    ablation=summary[summary.policy.isin(['fine_tuned','no_sensor_ids','no_attention','no_sensor_mask','no_temporal_mask'])&(summary.window=='early')&(summary.design=='farthest')] if len(summary) else summary
    if len(ablation):ablation=ablation[ablation.budget==ablation.groupby('dataset').budget.transform('max')]
    w.table('10_seed_runs/summary_raw_response.csv',ablation,skeys,[('gas','Gas'),('policy','Variant'),('mae_mean','MAE mean'),('mae_sd','MAE SD')],'Table 5. Ablation (ppm).')
    window=read('window_sensitivity/window_position_results.csv')
    if len(window):window=window[window.budget==window.budget.max()]
    w.table('window_sensitivity/window_position_results.csv',window,skeys,[('gas','Gas'),('window','Window'),('policy','Training'),('mae_mean','MAE mean'),('mae_sd','MAE SD')],'Table 6. Window-position sensitivity (ppm).')
    w.add('### V.G. Long-Term Drift\n\nTable 7 gives per-gas errors and normalized errors for the 20-label historical target transform, '
          'including the constant fitting-mean control. All budgets, macro errors, and loss/scaling comparisons are supplied as CSVs. '
          'The loss/scaling study is a sensitivity analysis, not test-based selection of a replacement primary model. A mean difference alone is insufficient to establish useful transfer under severe drift.')
    drift=read('10_seed_runs/summary_sensor_drift.csv')
    dkeys=['fold','budget','transform','policy','gas']
    main=drift[(drift.budget==20)&(drift['transform']=='log_huber')] if len(drift) else drift
    w.table('10_seed_runs/summary_sensor_drift.csv',main,dkeys,[('fold','Partition'),('gas','Gas'),('policy','Policy'),
            ('mae_mean','MAE mean'),('mae_sd','MAE SD'),('nmae_mean','NMAE mean'),('nmae_sd','NMAE SD')],'Table 7. Oracle-class conditional drift regression.')
    part10=read('sensor_drift/partition10_statistics.csv')
    macro10=part10[(part10.gas=='MacroMAE')&(part10.budget==20)&(part10['transform']=='log_huber')] if len(part10) else part10
    w.table('sensor_drift/partition10_statistics.csv',macro10,
            ['dataset','budget','design','window','gas','comparator','policy','transform','partition'],
            [('policy','Policy'),('improvement','MacroMAE difference'),('ci_low','95% CI lower'),('ci_high','95% CI upper')],'Table 8. Partition 10 paired seed uncertainty (ppm).')
    if len(macro10):
        r=macro10[macro10.policy=='fine_tuned'].iloc[0]
        w.add('For partition 10, the paired interval '+('includes zero: **no clear improvement**.' if r.ci_low<=0<=r.ci_high else
              'is below zero: fine-tuning increases error in this analysis.' if r.ci_high<0 else
              'is above zero in this seed-conditional comparison. This does not establish broad robustness to severe drift or practical significance.'))
    w.add('### V.H. Tail Error, Model Size, and Uncertainty\n\nQ0.9 is reported for scratch, frozen, and fine-tuned policies in the accompanying statistics CSV, '
          'alongside total, encoder, head, and trainable parameter counts and parameters per fitting exposure. Five overlapping windows do not multiply the effective independent supervised sample count. '
          'CO remains a particularly important quantitative limitation: evaluate its R2 and tail error alongside MAE, rather than equating an average relative gain with adequate calibration accuracy.\n\n'
          'Standard split-conformal coverage relies on exchangeability between calibration and test residuals. The chronological/session shift in this study violates that assumption. '
          'Intervals from the earlier split-conformal analysis are therefore diagnostic rather than guaranteed coverage intervals under deployment shift. '
          'The revision does not present the earlier coverage percentages as new 10-seed measurements; shift-aware methods [24] [25] remain future work.')
    tail=fixed[fixed.budget==fixed.groupby('dataset').budget.transform('max')] if len(fixed) else fixed
    w.table('10_seed_runs/summary_raw_response.csv',tail,skeys,
            [('gas','Gas'),('policy','Policy'),('r2_mean','Mean R2'),('rmse_mean','Mean RMSE'),('q90_mean','Mean Q0.9')],
            'Table 9. Full-budget quantitative and tail metrics (errors in ppm).')
    sizes=read('statistics/model_size.csv')
    if len(sizes):
        sizes=sizes[(sizes.dataset!='sensor_drift')&(sizes.seed==42)&(sizes.fold==0)&
                    (sizes.design=='farthest')&(sizes.window=='early')&sizes.policy.isin(['scratch','frozen','fine_tuned','generic_ssl'])]
        sizes=sizes[(sizes.budget==5)|(sizes.budget==sizes.groupby('dataset').budget.transform('max'))]
    w.table('statistics/model_size.csv',sizes,[k for k in skeys if k!='gas']+['seed'],
            [('dataset','Dataset'),('budget','Pairs'),('policy','Model'),('total','Total parameters'),('trainable','Trainable parameters'),
             ('encoder','Encoder'),('head','Head'),('fitting_exposures','Fitting exposures'),('parameters_per_fitting_exposure','Parameters/exposure')],
            'Table 10. Model size and effective supervised sample count; seed 42 exemplifies the fixed architecture and holdout sizes.')
    w.add('## VI. Discussion and Limitations\n\nThe primary evidence is the multi-fold comparison with a single predefined fine-tuned policy. '
          'Single-split improvements and favorable target-specific frozen scores must not replace that evidence. The added generic SSL control tests whether the particular '
          'sensor-aware architecture is useful beyond masked patch learning; it does not isolate sensor identity alone because the encoders differ. The ablation supplies the narrower component control. '
          'Low-budget outcomes must be interpreted jointly with calibration design and the number of actual fitting exposures.\n\n'
          'Physical sensor embeddings are tied to the array layout. These experiments support evaluation of same-system recalibration only and do not demonstrate cross-device transfer. '
          'A valid future cross-device experiment would require compatible arrays, meaningful channel correspondence, and a documented embedding-remapping or reinitialization strategy. '
          'Missing hardware/environment records restrict mechanistic interpretation. The same limitation applies to larger apparent H2 identity effects relative to NH3: '
          'sampling variability is an alternative explanation.\n\n'
          'Severe drift can erase or reverse transfer advantages. Partition-specific intervals and the fitting-mean control should govern interpretation, especially in partition 10. '
          'Domain alignment and shift-aware uncertainty are hypotheses for future evaluation, not demonstrated remedies in this experiment. '
          'The work measures requirements for labels used in adaptation, not savings in physical gas generation.')
    original=Document(next((ROOT/'revision_artifacts/environment').glob('*.docx')))
    refs=[p.text.replace('\ufffd','"') for p in original.paragraphs if re.match(r'^\[\d+\]',p.text)]
    w.add('## References\n\n'+'\n\n'.join(refs+[f'[{i}] {text}' for i,text in enumerate(ADDITIONS,18)]))
    path=out/'revised_manuscript.md';path.write_text('\n\n'.join(w.parts)+'\n',encoding='utf-8')
    dump(path.with_suffix('.claims.json'),w.claims)
    # Accessible review DOCX; journal typesetting remains a separate author step.
    doc=Document()
    for block in w.parts:
        lines=block.splitlines()
        table_lines=[line for line in lines if line.startswith('|')]
        if table_lines:
            for line in lines:
                if line and not line.startswith('|'):doc.add_paragraph(line)
            values=[[cell.strip() for cell in line.strip('|').split('|')] for line in table_lines]
            table=doc.add_table(rows=1,cols=len(values[0]));table.style='Table Grid'
            for cell,text in zip(table.rows[0].cells,values[0]):cell.text=text
            for row in values[2:]:
                for cell,text in zip(table.add_row().cells,row):cell.text=text
            continue
        for line in lines:
            if line.startswith('#'):
                level=len(line)-len(line.lstrip('#'));doc.add_heading(line.lstrip('# '),min(level,3))
            else:doc.add_paragraph(line)
    figures(run,out)
    from docx.shared import Inches
    doc.add_page_break();doc.add_heading('Figures for review',level=1)
    for name in ['fig2_preserved.png','fig3_pretraining.png']:
        if (out/name).exists():
            doc.add_picture(str(out/name),width=Inches(6.4));doc.add_paragraph(name)
    doc.save(out/'revised_manuscript_review.docx')
    response(run,w,complete)
    print(path)


def response(run,w,complete):
    out=run/'response';out.mkdir(exist_ok=True)
    comments=[
      ('Major 1','The premise is not tested; drop language implying reduced calibration cost.',
       'We restrict the claim to reduced labeled-calibration requirements. Existing labels are withheld during SSL; reduced exposure-generation cost is not demonstrated.', 'Abstract; I; VI','Scope correction; no cost experiment is claimed.'),
      ('Major 2','Pre-register one adaptation policy or select it using adaptation validation.',
       'Fine-tuning is predefined for every target. All policies share the same adaptation holdout, and scratch/frozen/fine-tuned results are retained separately.', 'III; V.B; Tables 1–2','See Table 1 and policy_selection/primary_policy_results.csv.'),
      ('Major 3','Use at least 10 seeds and error bars on budget and ablation results.',
       'The revised protocol uses independent pretraining and adaptation for all ten seeds; budget, baseline, and ablation tables report seed SD.', 'IV; V.C; V.F; Tables 2, 4–5','See tables and protocol.json; partial runs are not final evidence.'),
      ('Major 4','Promote multi-fold validation to the primary result.',
       'Six held-partition NH3/H2 folds and three rolling chronological CO/ethylene folds lead the Results. The paired bootstrap resamples seed/fold effects instead of pooling repeated predictions.', 'Abstract; V.B; Table 1','Table 1 supplies differences, intervals, relative changes, across-fold SD, and paired counts.'),
      ('Major 5','Partition 10 may be a null; add a fitting-mean constant predictor.',
       'We add a constant predictor fitted only on the adaptation fitting labels, per-gas and macro normalized errors, and paired partition-10 intervals. Wording follows the interval rather than a small mean difference.', 'V.G; Tables 7–8','See Table 8 and sensor_drift/partition10_statistics.csv.'),
      ('Major 6','Add a sensor-agnostic SSL and a chemometric baseline; define tree features.',
       'We implement a compact PatchTST-style masked channel-independent Transformer without sensor-ID embeddings, plus adaptation-tuned PLSR and RBF SVR. Trees are rerun with explicit flattened-window inputs.', 'II; V.E; Table 4','See Table 4, ssl_baselines/proposed_vs_generic_ssl.csv, and baseline feature documentation.'),
      ('Major 7','The selection is calibration design; add random selection.',
       'We use the term farthest-point calibration design and cross random/farthest designs with scratch/fine-tuned training over all fixed-split budgets.', 'III; V.D; Table 3','See Table 3 and calibration_design/paired_design_statistics.csv.'),
      ('Major 8','State sensor-ID array dependence and cross-device limitation.',
       'We explicitly limit scope to same-system transfer. No compatible cross-device acquisition was established, so we do not force an invalid array-remapping experiment.', 'VI','No cross-device result is claimed.'),
      ('Major 9','Characterize the private array, environmental conditions, and window placement.',
       'We document every unavailable hardware/environment field without inventing values. Four 64-sample window definitions are evaluated. Laboratory metadata and acquisition independence still require author-supplied records.', 'IV; V.F; Table 6; hardware/private_array_metadata.md','Table 6 gives sensitivity; hardware completeness remains unresolved.'),
      ('Moderate: CO','Acknowledge weak CO quantitative performance.',
       'We retain CO R2 and tail error and explicitly distinguish average relative improvement from adequate quantitative accuracy.', 'V.H','See the raw-response summary and statistics/q90_all_policies.csv.'),
      ('Moderate: conformal','State the exchangeability violation or use a shift-aware method.',
       'We explicitly state that chronological/session shift violates exchangeability and interpret earlier intervals as diagnostic. Earlier coverage percentages are not presented as rerun evidence.', 'V.H','No new guaranteed coverage result is claimed.'),
      ('Moderate: component variability','Explain sensor-identity variability and add error bars.',
       'The component comparison includes seed SD; gas-dependent effects cannot distinguish sensing selectivity from sampling variability.', 'V.F; VI; Table 5','See Table 5.'),
      ('Moderate: normalization and target loss','Report normalized errors and compare log-target versus linear-target objectives.',
       'Per-gas and macro NMAE accompany MAE. The 20-label control compares historical log-Huber, fixed-log Huber, fixed-linear Huber, and fixed-linear MAE; tiny budgets use fixed scaling.', 'IV; V.G','See sensor_drift_loss/target_transform_comparison.csv and scaling_sensitivity.csv.'),
      ('Moderate: size and repeated exposures','Report parameter counts, effective samples, repeated exposures, and scratch Q0.9.',
       'Every neural job records parameter counts and fitting exposure counts, and every policy records Q0.9. Exposure IDs and signal fingerprints are audited; physical independence remains unverified without logs.', 'IV; V.H','623 exposure/window fingerprints audited; zero identical fingerprints and zero ID overlap across 11 split definitions.'),
      ('Moderate: competing methods and literature','Reimplement a competing method or published split; expand references.',
       'The PatchTST-style reimplementation is described with its deviations from the published forecasting setup. Ten relevant sources extend the bibliography to 27 references.', 'II; V.E; References','The comparison is a matched gas-regression implementation, not a claim of reproducing forecasting scores.'),
      ('Presentation','Correct wording, keywords, duplicate tables, and unreadable Fig. 3.',
       'The rewritten review draft uses a single all-policy budget table, corrected language and keywords, and vector reconstruction figures with readable fonts; the supplied Fig. 2 is preserved.', 'Abstract; V; Figure 3','Final IEEE two-column typesetting and visual proof review remain an author production step.'),
    ]
    lines=['# Point-by-point reviewer response','',
           'Status: '+('full rerun artifacts available; hardware metadata and journal layout still require author review.' if complete else 'interim; experiments are still incomplete, so numerical claims are not final.'),'',
           'Comment summaries below paraphrase the supplied reviewer PDF. Original text is preserved in the environment snapshot.','']
    for title,comment,answer,location,result in comments:
        lines += [f'### {title}','',f'> {comment}','','**Response:**',answer,'','**Changes made:**',
                  f'- {answer}','','**New result / status:**',result,'','**Location in manuscript:**',location,'']
    # Embed exact generated numerical tables in the response so major comments
    # have reviewable numbers rather than merely future-tense promises.
    lines += ['## Numerical evidence for major comments','']
    for part in w.parts:
        if part.startswith('Table '):lines += [part,'']
    (out/'reviewer_response.md').write_text('\n'.join(lines),encoding='utf-8')


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--run-id',required=True);p.add_argument('--allow-incomplete',action='store_true')
    a=p.parse_args();build(ROOT/'revision_artifacts/runs'/a.run_id,a.allow_incomplete)
