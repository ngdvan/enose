"""Minimal edits to the original IEEE DOCX; matching clean/highlighted copies.

Uses the frozen evidence snapshot, not incomplete jobs or the broad rewrite.
"""
from copy import deepcopy
from difflib import SequenceMatcher
import hashlib
import json
import re
import sys
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from lxml import etree
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'code'))
from revision_common import ROOT, dump, sha256

OUT=ROOT/'author_review/minimal_revision_001'
EVIDENCE=OUT/'evidence'
SOURCE=next((ROOT/'revision_artifacts/environment').glob('*.docx'))
RUN=ROOT/'revision_artifacts/runs/major_revision_001'
LEGACY=ROOT/'revision_artifacts/environment/legacy_results'
changes=[]


def replace(p,new,location,highlight=True):
    """Token diff preserving formatting on unchanged text and anchored drawings."""
    old=p.text
    if old==new:return
    assert not p._p.xpath('.//m:oMath'), f'Do not flatten equation paragraph: {location}'
    oldruns=[];pos=0
    for r in p.runs:
        if r.text:oldruns.append((pos,pos+len(r.text),r.text,deepcopy(r._r.rPr)))
        pos+=len(r.text)
    complex_runs=[deepcopy(r._r) for r in p.runs if r._r.xpath('.//w:drawing|.//w:pict')]
    for child in list(p._p):
        if child.tag!=qn('w:pPr'):p._p.remove(child)
    tokens=lambda s:re.findall(r'\w+|[^\w\s]+|\s+',s)
    a,b=tokens(old),tokens(new)
    offsets=np.cumsum([0]+[len(x) for x in a]).tolist()
    def append(text,properties,mark):
        if not text:return
        run=p.add_run(text)
        if properties is not None:run._r.insert(0,deepcopy(properties))
        if mark and highlight:run.font.highlight_color=WD_COLOR_INDEX.YELLOW
    for tag,i,j,k,l in SequenceMatcher(None,a,b,autojunk=False).get_opcodes():
        if tag=='equal':
            start,end=offsets[i],offsets[j]
            for lo,hi,txt,pr in oldruns:
                if hi>start and lo<end:append(txt[max(start-lo,0):min(end-lo,len(txt))],pr,False)
        elif tag in ['insert','replace']:
            props=next((pr for lo,hi,txt,pr in oldruns if hi>offsets[i]),oldruns[0][3] if oldruns else None)
            append(''.join(b[k:l]),props,True)
    for r in complex_runs:p._p.append(r)
    assert p.text==new,(location,p.text,new)
    changes.append(dict(location=location,kind='text revision',before=old,after=new))


def insert_after(p,text,location):
    node=OxmlElement('w:p')
    if p._p.pPr is not None:node.append(deepcopy(p._p.pPr))
    p._p.addnext(node);new=Paragraph(node,p._parent)
    replace(new,text,location)
    return new


def replace_literal(p,old,new,location):
    """Replace a plain phrase without touching interleaved OMML equations."""
    for run in p.runs:
        if old in run.text:
            run.text=run.text.replace(old,new);run.font.highlight_color=WD_COLOR_INDEX.YELLOW
            changes.append(dict(location=location,kind='phrase revision',before=old,after=new));return
    raise ValueError(f'Phrase crosses runs: {location}: {old}')


def delete(node,location,reason):
    changes.append(dict(location=location,kind='deletion',reason=reason,
                        before=''.join(t.text or '' for t in node.iter(qn('w:t')))))
    node.getparent().remove(node)


def move_block(first,last,before,location):
    nodes=[];node=first
    while node is not last:
        nodes.append(node);node=node.getnext()
    for node in nodes:before.addprevious(node)
    for run in Paragraph(first,None).runs:run.font.highlight_color=WD_COLOR_INDEX.TURQUOISE
    changes.append(dict(location=location,kind='section moved',reason='Place the primary generalization analysis before fixed-split results, and baselines before ablations.'))


def cell_text(cell,text,location):
    replace(cell.paragraphs[0],text,location)
    for p in cell.paragraphs[1:]:delete(p._p,location,'Remove redundant cell paragraph after updating the value.')


def configure_table(table,widths):
    table.autofit=False
    for column,width in zip(table.columns,widths):column.width=Inches(width)
    for row in table.rows:
        # Let rows wrap naturally, but do not split individual rows across pages.
        trpr=row._tr.get_or_add_trPr()
        if trpr.find(qn('w:cantSplit')) is None:trpr.append(OxmlElement('w:cantSplit'))
        for cell,width in zip(row.cells,widths):
            cell.width=Inches(width)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.space_before=Pt(2)
                for run in p.runs:run.font.size=Pt(8)
    header=table.rows[0]._tr.get_or_add_trPr()
    header.append(OxmlElement('w:tblHeader'))


def make_figures(summary,pred):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':8,'axes.labelsize':8,
                         'axes.titlesize':9,'legend.fontsize':7,'pdf.fonttype':42,'svg.fonttype':'none'})
    fig,axes=plt.subplots(3,1,figsize=(3.45,3.9))
    for ax,dataset,title in zip(axes,['nh3_h2','co_ethylene','sensor_drift'],['NH3/H2','CO/ethylene','Sensor_Drift']):
        if dataset=='sensor_drift':
            j=json.loads((LEGACY/'sensor_drift/results.json').read_text())
            curves=list(j['pretraining_histories'].values())
        else:
            curves=[json.loads(p.read_text())['loss'] for p in sorted((RUN/'checkpoints').glob(f'{dataset}_f0_early_s*_full.json'))]
            assert len(curves)==10
        a=np.asarray(curves);m=a.mean(0);sd=a.std(0,ddof=1);epochs=np.arange(1,len(m)+1)
        ax.plot(epochs,m,color='#24557d',lw=1.1);ax.fill_between(epochs,m-sd,m+sd,color='#24557d',alpha=.2)
        ax.set(title=f'{title} ({len(a)} repeats)',ylabel='Loss',xlabel='Epoch')
        ax.grid(alpha=.15)
    fig.tight_layout(pad=.6,h_pad=.65)
    for ext in ['png','pdf','svg']:fig.savefig(OUT/f'fig3_revised.{ext}',dpi=400)
    plt.close(fig)
    fig,axes=plt.subplots(2,2,figsize=(3.45,3.45))
    colors={'scratch':'#666666','frozen':'#287a9c','fine_tuned':'#c85a2c'}
    for ax,gas in zip(axes.flat,['NH3','H2','CO','Ethylene']):
        d=pred[pred.gas==gas];d=d[d.budget==d.budget.max()]
        limit=20 if gas=='Ethylene' else 533.33 if gas=='CO' else 500
        for policy in ['scratch','frozen','fine_tuned']:
            p=d[d.policy==policy];ax.scatter(p.true,p.prediction,s=3,alpha=.2,color=colors[policy],label=policy.replace('_',' '),rasterized=True)
        ax.plot([0,limit],[0,limit],color='black',lw=.65,ls='--')
        ax.set(title=gas,xlim=(0,limit),ylim=(0,limit),xlabel='Reference (ppm)',ylabel='Predicted (ppm)')
        ax.tick_params(labelsize=7)
    handles,labels=axes[0,0].get_legend_handles_labels()
    fig.legend(handles,labels,loc='lower center',ncol=3,fontsize=6.5,markerscale=2)
    fig.tight_layout(pad=.5,rect=[0,.08,1,1])
    for ext in ['png','pdf','svg']:fig.savefig(OUT/f'fig4_revised.{ext}',dpi=400)
    plt.close(fig)


def replace_figure(p,path,width,location):
    for r in list(p.runs):
        if r._r.xpath('.//w:drawing'):p._p.remove(r._r)
    p.add_run().add_picture(str(path),width=Inches(width))
    changes.append(dict(location=location,kind='figure replaced',after=path.name,
                        reason='Use completed evidence and readable labels; image changes are identified by the highlighted caption.'))


def main():
    original_hash=sha256(SOURCE)
    d=Document(SOURCE);paragraphs=list(d.paragraphs);tables=list(d.tables)
    # Preserve the actual IEEE document, including OMML, styles, authors and anchors.
    summary=pd.read_csv(EVIDENCE/'completed_10seed_summary.csv')
    pred=pd.read_csv(EVIDENCE/'completed_10seed_predictions.csv')
    multi=pd.read_csv(EVIDENCE/'historical_multifold_recomputed_statistics.csv')
    ablation=pd.read_csv(EVIDENCE/'historical_ablation_with_sd.csv')
    drift=pd.read_csv(EVIDENCE/'historical_drift_macro_summary.csv')
    gasdrift=pd.read_csv(EVIDENCE/'historical_drift_gas_summary.csv')
    sizes=pd.read_csv(EVIDENCE/'model_size.csv')
    def result(gas,policy,budget=None):
        a=summary[(summary.gas==gas)&(summary.policy==policy)]
        return a[a.budget==(budget or a.budget.max())].iloc[0]
    def fmt(x,gas='NH3'):return f'{float(x):.{3 if gas=="Ethylene" else 2}f}'
    def ms(row,gas='NH3'):return f'{fmt(row.mae,gas)} ± {fmt(row.sd,gas)}'
    def gain(gas,budget=None):return 100*(1-result(gas,'fine_tuned',budget).mae/result(gas,'scratch',budget).mae)
    make_figures(summary,pred)
    abstract=next(p for row in tables[0].rows for cell in row.cells for p in cell.paragraphs if p.text.startswith('Abstract'))
    newabstract=('Abstract— Electronic nose calibration usually relies on controlled gas exposures with known concentrations. '
       'The present experiments reuse controlled exposures with their concentration labels withheld during pretraining, so they examine reduced labeled-calibration requirements rather than reduced gas-generation cost. '
       'This study explores sensor-aware self-supervised learning through a shared dilated one-dimensional convolutional encoder, sensor-identity embeddings, cross-sensor attention, and reconstruction of masked time spans and complete sensor channels. '
       'The resulting encoder is adapted using a limited number of concentration labels. Fine-tuning is the predefined primary transfer policy for all targets. '
       'The evaluation includes a private 20-channel NH3/H2 array, a public 16-channel CO/ethylene array, and a public 36-month sensor drift dataset. '
       'The primary generalization analysis uses six held-partition NH3/H2 folds and three rolling CO/ethylene folds with three repeats. '
       'Mean relative MAE reductions are 16.4%, 20.1%, 20.4%, and 14.5% for NH3, H2, CO, and ethylene, respectively; seed-and-fold bootstrap intervals include zero except for H2. '
       'Secondary fixed-split label-budget and classical-baseline comparisons use ten seeds. Classical models outperform the proposed model on some targets, and budget effects are not uniformly monotonic. '
       'Sensor drift regression is oracle-class conditional; the small partition-10 mean difference does not establish robust improvement under severe drift. '
       'The evidence concerns same-system recalibration and does not establish cross-device transfer.')
    replace(abstract,newabstract,'Abstract')
    replacements={
      6:paragraphs[6].text.replace('Collecting these labels becomes costly','Generating these controlled exposures becomes costly'),
      7:'In this study, responses from previously controlled exposures are reused with their concentration labels withheld during pretraining. They retain temporal response shapes and recurring relationships among sensors. The experiment tests whether these patterns can reduce the labeled-calibration requirements of a later adaptation stage; it does not test cheaper generation of gas exposures.',
      10:paragraphs[10].text.replace('Transfer improves two-gas regression and oracle-class gas-conditional concentration estimates under drift.','The experiments distinguish fixed-split label efficiency from multi-fold generalization and oracle-class conditional concentration regression under drift.'),
      17:paragraphs[17].text.replace('reduce the labeled-calibration effort','reduce labeled-calibration requirements'),
      25:paragraphs[25].text.replace('a larger set of unlabeled exposures is available','a larger set of previously acquired exposures is available with its concentration labels withheld'),
      41:paragraphs[41].text+' These embeddings are tied to the physical array layout. The current experiments concern same-system transfer and do not demonstrate cross-device transfer.',
      67:'Adaptation Policies and Calibration-Pair Design',
      68:paragraphs[68].text.replace('To integrate pretrain result with labeled data, three adaptation policies are selected','To integrate pretraining with labeled data, three adaptation policies are evaluated')+' Fine-tuning is the predefined primary policy for every target; frozen transfer is reported separately and is not selected using test MAE.',
      74:paragraphs[74].text.replace('0�500','0–500').replace('six repeated exposures','six additional exposure identifiers at repeated concentration coordinates')+' These partitions are inferred from consecutive Combination identifiers, not independently verified acquisition days. An exposure/window fingerprint audit found no identical signals; distinct identifiers alone do not prove independent physical acquisition.',
      104:'The reconstruction curves in Fig. 3 summarize ten independently pretrained raw-response models per dataset and the existing three Sensor_Drift repeats. The shaded bands show one standard deviation across repeats. Reconstruction is an optimization diagnostic; its numerical scale differs between raw responses and drift descriptors.',
      107:'Fig. 3. Self-supervised reconstruction loss: mean ± SD across ten raw-response repeats and three existing Sensor_Drift repeats.',
      108:'Fixed-Split Label-Budget Analysis',
      109:'Table I reports scratch, frozen-transfer, and fine-tuned-transfer MAE over ten seeds. These fixed-split results are a controlled analysis of labeled-calibration requirements; the multi-fold analysis above remains the primary generalization evidence.',
      110:'At the full label budget, predefined fine-tuning gives '+', '.join(f'{gas} MAE of {fmt(result(gas,"fine_tuned").mae,gas)} ppm ({gain(gas):.1f}% lower than scratch)' for gas in ['NH3','H2','CO','Ethylene'])+'. Fig. 4 compares the predictions of all three policies on the held-out exposures.',
      113:'Fig. 4. Full-budget predicted versus reference concentrations for scratch, frozen transfer, and fine-tuning. Each point is a held-out exposure prediction from one of the ten seeds; the dashed line is identity.',
      115:'At five and ten labeled pairs, predefined fine-tuning reduces CO MAE relative to scratch by '+f'{gain("CO",5):.1f}% and {gain("CO",10):.1f}%, respectively. The corresponding H2 changes are {gain("H2",5):.1f}% and {gain("H2",10):.1f}%. These results are specific to the fixed split and coverage-based calibration-pair design.',
      116:'Table I gives mean ± SD over ten seeds. Pair budgets count all labeled adaptation pairs, including the group-disjoint validation subset. Neither fine-tuning superiority over frozen transfer nor monotonic improvement with increasing budget is assumed. All policies use the same fitting and validation exposure groups within each seed and budget.',
      117:'Table I. Fixed-Split Calibration MAE (ppm), Mean ± SD Over Ten Seeds',
      119:'Primary Multi-Fold Validation',
      123:'The ablation analysis indicates gas-dependent component effects. Table III reports mean ± SD of paired MAE changes across the existing three repeats. These estimates describe this particular ablation protocol and do not establish a single component as uniformly beneficial.',
      124:'Positive values indicate higher error after component removal. The ablation protocol reserves 20% of adaptation groups for validation. The larger sensor-identity effect for H2 than NH3 could reflect channel selectivity, sampling variability, or both; the current experiments cannot separate these explanations.',
      125:'Table III. Ablation MAE Changes (ppm), Mean ± SD Across Three Existing Repeats',
      126:'Baseline Comparison',
      127:'Table II compares ten-seed full-budget results using the same adaptation/test exposures and group-disjoint fitting/validation subsets. Extra Trees and XGBoost are rerun using flattened, pretraining-scaled 64-sample windows (1,280 NH3/H2 or 1,024 CO/ethylene features). PLSR [18] and separate-target RBF SVR [19] additionally standardize features using fitting data only. Components for PLSR and C, gamma, and epsilon for SVR are chosen by adaptation-validation mean normalized MAE, without test-based tuning. The proposed model is not uniformly strongest: PLSR has lower CO MAE, Extra Trees has lower ethylene MAE, and SVR is close to fine-tuning on H2.',
      128:'Table II. Full-Budget Baseline MAE (ppm), Mean ± SD Over Ten Seeds',
      130:'The Sensor_Drift results in Table IV retain the existing three-repeat oracle-class conditional experiment. Fine-tuning gives lower mean macro-MAE than scratch in partitions 8 and 9. Partition 10 has much larger absolute error; the small mean difference does not establish robust improvement under severe drift. The constant control predicts the per-gas mean of the adaptation fitting subset used by the historical fine-tuned model, excluding validation and test labels. The historical neural policies used policy-dependent validation splits, which limits matched-policy interpretation. Table V reports per-gas normalized error for fine-tuning; complete per-policy MAE and NMAE are supplied in the companion evidence CSVs. NMAE divides MAE by each gas concentration range in historical partitions 1–6, using those reference labels only for reporting, not for pretraining, fitting, or model selection. A fixed reference is necessary because some test gas/partition combinations have zero concentration range; macro-NMAE averages all six gases.',
      132:'Table IV. Sensor_Drift Macro-MAE (ppm), Mean ± SD Across Three Existing Repeats; Matched Constant Control',
      134:'At full budget, mean Q0.9 absolute errors for scratch, frozen transfer, and fine-tuning are '+ '; '.join(f'{gas}: '+', '.join(fmt(result(gas,p).q90,gas) for p in ['scratch','frozen','fine_tuned'])+' ppm' for gas in ['NH3','H2','CO','Ethylene'])+'. CO remains quantitatively weak: fine-tuning has mean R2 '+f'{result("CO","fine_tuned").r2:.3f} and mean Q0.9 {result("CO","fine_tuned").q90:.2f} ppm on the 0–533.33 ppm scale.',
      135:paragraphs[135].text.replace('We also use split-conformal prediction as a diagnostic.','The existing three-repeat analysis also used split-conformal prediction as a diagnostic; these are not new ten-seed coverage estimates.'),
      136:'Standard split-conformal coverage relies on exchangeability between calibration and test residuals. The chronological/session shift in this study violates that assumption. The reported intervals are therefore diagnostic rather than guaranteed coverage intervals under deployment shift [20]. Shift-aware calibration requires separate evaluation before deployment.',
      138:'The ablations suggest that channel identity and structured masking can contribute differently across gases. They do not establish superiority over a sensor-agnostic self-supervised method, and the larger identity effect for H2 than NH3 cannot be separated from sampling variability or unrecorded sensor characteristics. The classical-baseline results also show that a transferred neural representation is not necessarily the strongest regressor for every gas.',
      139:'The fixed-split experiments show reductions in labeled-calibration requirements relative to scratch under the predefined policy and coverage-based calibration design. They do not demonstrate reduced gas-generation cost, and label-budget effects are not uniformly monotonic. Known concentration coordinates are used in the experimental design; this is not generic acquisition of unknown labels.',
      140:paragraphs[140].text.replace('An explicit temporal- or domain-alignment objective added after sensor-aware pretraining could improve robustness under stronger drift.','Explicit temporal or domain alignment is a possible future direction for severe drift; its benefit is not established by these experiments.'),
      142:paragraphs[142].text.replace('unlabeled Electronic nose measurements','previously acquired electronic nose responses with withheld concentration labels'),
      143:'Multi-fold validation is the primary generalization analysis. Its mean differences favor fine-tuning, but the seed-and-fold bootstrap intervals include zero for NH3, CO, and ethylene in the existing three-repeat evidence. The ten-seed fixed-split comparisons are secondary and do not establish universal superiority over classical models.',
      144:'On the 36-month Sensor_Drift dataset, the existing three-repeat results suggest useful transfer in partitions 8–9 but do not establish robust improvement in partition 10. These results are conditional on known gas identity. Sensor-ID embeddings bind the encoder to the array layout; current evidence concerns same-system recalibration and does not demonstrate cross-device transfer.',
    }
    for idx,text in replacements.items():replace(paragraphs[idx],text,f'Original paragraph {idx}')
    # Leave equations, math symbols, and their surrounding runs intact.
    replace_literal(paragraphs[69],'The calibration pairs are selected to cover the available concentration space by farthest-point sampling,',
        'Calibration pairs are designed using known concentration coordinates to cover the available space by farthest-point sampling,','Calibration design preceding Eq. (16)')
    # Remove the floating test-selected summary, including its legacy fallback.
    for r in list(paragraphs[92].runs):
        if r._r.xpath('.//w:drawing'):
            delete(r._r,'Original floating Table I','Remove the test-selected policy summary and duplicated full-budget metrics.')
    # In this source the VML fallback is held in AlternateContent at paragraph level.
    for node in list(paragraphs[92]._p):
        if 'Summary of Full-Budget Regression Performance' in ''.join(t.text or '' for t in node.iter(qn('w:t'))):
            delete(node,'Original floating Table I fallback','Remove duplicate drawing fallback of the same obsolete summary.')
    delete(tables[1]._tbl,'Original Table II','Remove the duplicate full-budget-only table; Table I contains all policies and budgets with uncertainty.')
    delete(paragraphs[111]._p,'Original Table II caption','Caption belongs to the removed duplicate table.')
    delete(paragraphs[114]._p,'Original Label-Efficient Calibration heading','Merge with the fixed-split label-budget subsection, without changing the method.')
    insert_after(paragraphs[74],
        'Authoritative records for sensor manufacturer/model, sensing material, heater temperature, acquisition circuit, sampling rate, chamber volume, gas source, flow controller/rate, carrier gas, baseline convention, exposure/purge duration, ambient temperature, relative humidity/control, and reference concentration accuracy were unavailable in the inspected material. These omissions limit reproducibility and physical interpretation. The 0–500 ppm setting is a controlled calibration benchmark, not validation of low-concentration ambient monitoring.',
        'Private-array characterization limitation')
    insert_after(paragraphs[91],
        'The revised fixed-split raw-response and classical-baseline experiments use ten seeds (42, 142, 242, 342, 442, 542, 642, 742, 842, 942), with pretraining repeated for each seed. A common 20% exposure-group adaptation holdout selects checkpoints and classical hyperparameters. The multi-fold, ablation, drift, and conformal results retained below use the existing three repeats and are identified separately.',
        'Repeat counts and common validation protocol')
    for idx,dataset in [(120,'nh3_h2'),(121,'co_ethylene')]:
        gases=['NH3','H2'] if dataset=='nh3_h2' else ['CO','Ethylene']
        intro=('The primary NH3/H2 analysis uses six leave-one-partition-out folds. Each fold holds out one partition for testing, uses the preceding cyclic partition for adaptation, and uses the remaining four for pretraining. '
               'This is held-partition generalization, not uniformly forward chronological evaluation. The existing three repeats give 18 fold/repeat runs at 57 pairs. ' if dataset=='nh3_h2' else
               'The dynamic gas mixture analysis divides acquisition order into seven blocks and uses three rolling pretrain–adapt–test folds. The existing three repeats give nine fold/repeat runs at the common 20-pair budget. ')
        sentences=[]
        for gas in gases:
            r=multi[multi.gas==gas].iloc[0]
            sentences.append(f'{gas} mean MAE is {fmt(r.scratch_mae,gas)} ppm for scratch and {fmt(r.transfer_mae,gas)} ppm for fine-tuning, a paired mean reduction of {fmt(r.improvement,gas)} ppm ({r.relative_improvement_pct:.1f}%). '
                f'The recomputed 95% interval is [{fmt(r.ci_low,gas)}, {fmt(r.ci_high,gas)}] ppm; the across-fold SD of fine-tuned fold means is {fmt(r.across_fold_sd,gas)} ppm, with {int(r.n_paired_predictions)} paired predictions across repeats.')
        replace(paragraphs[idx],intro+' '.join(sentences),f'Primary multi-fold results: {dataset}')
    insert_after(paragraphs[121],
        'These intervals use 5,000 paired bootstrap replicates that resample seed and fold indices, weighting folds equally, rather than treating repeated predictions as independent exposures. They are conditional on the observed test exposures; overlapping training folds and few held-out blocks limit population inference. The historical runs also used policy-dependent adaptation holdouts. Intervals include zero for NH3, CO, and ethylene, so uniformly positive generalization evidence is not established.',
        'Corrected multi-fold uncertainty and historical protocol limitation')
    # Move the historical multi-fold block, including its newly inserted caveat.
    move_block(paragraphs[119]._p,paragraphs[122]._p,paragraphs[108]._p,'Multi-fold promoted to first quantitative Results subsection')
    move_block(paragraphs[126]._p,paragraphs[129]._p,paragraphs[122]._p,'Baseline comparison before ablation')
    # Update all budget cells; one main table replaces two redundant summaries.
    for i,row in enumerate(tables[2].rows[1:],1):
        gas=row.cells[0].text;budget=int(row.cells[1].text)
        for cell,policy in zip(row.cells[2:],['scratch','frozen','fine_tuned']):
            cell_text(cell,ms(result(gas,policy,budget),gas),f'Table I {gas}, {budget} pairs, {policy}')
    configure_table(tables[2],[.42,.34,.91,.91,.92])
    for row in tables[3].rows[1:]:
        label=row.cells[0].text
        variant={'None: full model':'full','Sensor-ID embedding':'no_sensor_ids','Cross-sensor attention':'no_attention',
                 'Temporal-span masking':'no_temporal_mask','Complete-sensor masking':'no_sensor_mask'}[label]
        for cell,gas in zip(row.cells[1:],['NH3','H2','CO','Ethylene']):
            r=ablation[(ablation.gas==gas)&(ablation.variant==variant)].iloc[0]
            cell_text(cell,f'{r.delta:+.2f} ± {r.sd:.2f}',f'Table III {variant}, {gas}')
    configure_table(tables[3],[1.1,.6,.6,.6,.6])
    # Six models by four targets is more readable than adding two narrow columns.
    oldbase=tables[4];table=d.add_table(rows=7,cols=5)
    table._tbl.remove(table._tbl.tblPr);table._tbl.insert(0,deepcopy(oldbase._tbl.tblPr))
    headers=['Model','NH3','H2','CO','Ethylene']
    for cell,text in zip(table.rows[0].cells,headers):cell_text(cell,text,'Table II header')
    for row,policy,label in zip(table.rows[1:],['scratch','extra_trees','xgboost','plsr','svr','fine_tuned'],['Scratch','Extra Trees','XGBoost','PLSR','SVR','Fine-tuned']):
        cell_text(row.cells[0],label,f'Table II {label}')
        for cell,gas in zip(row.cells[1:],['NH3','H2','CO','Ethylene']):cell_text(cell,ms(result(gas,policy),gas),f'Table II {label}, {gas}')
    oldbase._tbl.addprevious(table._tbl)
    delete(oldbase._tbl,'Original Table V','Replace with matched ten-seed comparison including PLSR/SVR and uncertainty; all six models share the declared input/holdout protocol.')
    configure_table(table,[.7,.7,.7,.7,.7])
    # Retain existing drift experiments with an honestly matched new control.
    for row,policy in zip(tables[5].rows[1:],['scratch','frozen','fine_tuned']):
        for cell,partition in zip(row.cells[1:],[8,9,10]):
            r=drift[(drift.partition==partition)&(drift.policy==policy)].iloc[0]
            cell_text(cell,ms(r)+' ppm',f'Table IV {policy}, partition {partition}')
    row=tables[5].add_row();cell_text(row.cells[0],'Constant mean','Table IV constant label')
    for cell,partition in zip(row.cells[1:],[8,9,10]):
        r=drift[(drift.partition==partition)&(drift.policy=='constant')].iloc[0]
        cell_text(cell,ms(r)+' ppm',f'Table IV constant, partition {partition}')
    configure_table(tables[5],[.85,.88,.88,.89])
    caption=Paragraph(OxmlElement('w:p'),d._body)
    caption.style=paragraphs[132].style
    replace(caption,'Table V. Fine-Tuned Per-Gas NMAE (%), Mean ± SD Across Three Existing Repeats','New normalized error caption')
    normtable=d.add_table(rows=8,cols=4)
    normtable._tbl.remove(normtable._tbl.tblPr);normtable._tbl.insert(0,deepcopy(tables[5]._tbl.tblPr))
    for cell,text in zip(normtable.rows[0].cells,['Gas','Partition 8','Partition 9','Partition 10']):cell_text(cell,text,'Table V header')
    for row,gas in zip(normtable.rows[1:7],drift_names()):
        cell_text(row.cells[0],gas,'Table V gas')
        for cell,partition in zip(row.cells[1:],[8,9,10]):
            r=gasdrift[(gasdrift.gas==gas)&(gasdrift.partition==partition)&(gasdrift.policy=='fine_tuned')].iloc[0]
            cell_text(cell,f'{100*r.nmae:.2f} ± {100*r.nmae_sd:.2f}',f'Table V {gas}, partition {partition}')
    cell_text(normtable.rows[7].cells[0],'Macro-NMAE','Table V macro')
    for cell,partition in zip(normtable.rows[7].cells[1:],[8,9,10]):
        r=drift[(drift.partition==partition)&(drift.policy=='fine_tuned')].iloc[0]
        cell_text(cell,f'{100*r.nmae:.2f} ± {100*r.nmae_sd:.2f}',f'Table V macro, partition {partition}')
    tables[5]._tbl.addnext(caption._p);caption._p.addnext(normtable._tbl)
    configure_table(normtable,[.98,.84,.84,.84])
    # Report capacity relative to independent fitting exposures, not window count.
    countlines=[]
    for dataset,label in [('nh3_h2','NH3/H2'),('co_ethylene','CO/ethylene')]:
        r=sizes[(sizes.dataset==dataset)&(sizes.policy=='fine_tuned')&(sizes.seed==42)&(sizes.budget==sizes[sizes.dataset==dataset].budget.max())].iloc[0]
        frozen=sizes[(sizes.dataset==dataset)&(sizes.policy=='frozen')&(sizes.seed==42)].iloc[0]
        countlines.append(f'{label}: {int(r.total):,} total parameters ({int(r.encoder):,} encoder, {int(r["head"]):,} head); '
                          f'{int(r.total):,} are trainable during fine-tuning and {int(frozen.trainable):,} during frozen transfer. '
                          f'The full-budget fitting subset has {int(r.fitting_exposures)} exposures, or approximately {r.parameters_per_fitting_exposure:,.0f} parameters per fitting exposure.')
    insert_after(paragraphs[134],' '.join(countlines)+' The five offset windows per exposure do not represent five independent supervised samples.','Parameter counts and effective supervised sample size')
    # Three necessary new citations for methods actually discussed in this copy.
    refs=[
       '[18] S. Wold, M. Sjöström, and L. Eriksson, “PLS-regression: a basic tool of chemometrics,” Chemometrics and Intelligent Laboratory Systems, vol. 58, pp. 109–130, 2001, doi: 10.1016/S0169-7439(01)00155-1.',
       '[19] A. J. Smola and B. Schölkopf, “A tutorial on support vector regression,” Statistics and Computing, vol. 14, pp. 199–222, 2004, doi: 10.1023/B:STCO.0000035301.49549.88.',
       '[20] R. F. Barber, E. J. Candès, A. Ramdas, and R. J. Tibshirani, “Conformal prediction beyond exchangeability,” Annals of Statistics, vol. 51, 2023, doi: 10.1214/23-AOS2276.',
    ]
    anchor=paragraphs[162]
    for ref in refs:anchor=insert_after(anchor,ref,'Required supporting reference')
    replace_figure(paragraphs[106],OUT/'fig3_revised.png',3.45,'Figure 3')
    replace_figure(paragraphs[112],OUT/'fig4_revised.png',3.45,'Figure 4')
    for idx in [107,113]:
        for run in paragraphs[idx].runs:run.font.highlight_color=WD_COLOR_INDEX.YELLOW
    # The original source already has valid Index Terms; do not rewrite it.
    # Check all visible stories for leftover typo tokens and obsolete summary.
    alltext=' '.join(t.text or '' for t in d._element.iter(qn('w:t')))
    assert 'Best policy' not in alltext
    assert not re.search(r'\b(overal|trainning|comparision|Multif-fold)\b',alltext)
    highlighted=OUT/'IEEE_V2.2_Minimal_Revision_Highlighted.docx'
    clean=OUT/'IEEE_V2.2_Minimal_Revision_Clean.docx'
    d.save(highlighted)
    # Identical content, layout and figures; only highlight XML differs.
    for element in list(d._element.iter(qn('w:highlight'))):element.getparent().remove(element)
    d.save(clean)
    assert sha256(SOURCE)==original_hash
    dump(OUT/'change_manifest.json',changes)
    validate(clean,highlighted)
    print(clean);print(highlighted)


def drift_names():return ['Ethanol','Ethylene','Ammonia','Acetaldehyde','Acetone','Toluene']


def validate(clean,highlighted):
    a,b=Document(clean),Document(highlighted)
    count=len(b._element.xpath('.//w:highlight'))
    assert count>0 and not a._element.xpath('.//w:highlight')
    for x in list(b._element.xpath('.//w:highlight')):x.getparent().remove(x)
    assert etree.tostring(a._element)==etree.tostring(b._element),'Clean and highlighted content must match exactly'
    original=Document(SOURCE)
    def math_content(doc):
        # Exclude equations embedded inside the deliberately deleted Table I drawing.
        result=[]
        for m in doc._element.xpath('.//m:oMath'):
            if any(ancestor.tag in [qn('w:drawing'),qn('w:pict'),qn('w:txbxContent')] for ancestor in m.iterancestors()):continue
            result.append(etree.tostring(m,method='c14n'))
        return result
    assert math_content(original)==math_content(a),'Original scientific equations changed'
    assert [s._sectPr.xml for s in original.sections]==[s._sectPr.xml for s in a.sections],'Original page/column settings changed'
    dump(OUT/'document_validation.json',dict(clean_highlighted_text_and_structure_identical=True,
        clean_highlight_count=0,highlighted_runs=count,original_equation_objects_preserved=len(math_content(a)),
        original_section_settings_preserved=True,source_unchanged=True,
        original_sha256=sha256(SOURCE),clean_sha256=sha256(clean),highlighted_sha256=sha256(highlighted),
        publication_ready=False,reason='Several requested ten-seed analyses and authoritative hardware metadata are still outstanding.'))


if __name__=='__main__':main()
