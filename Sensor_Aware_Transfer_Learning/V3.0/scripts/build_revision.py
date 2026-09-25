"""Build V3.0 from recovered editable V2.2 sources, preserving tables/OMML/media."""
from pathlib import Path
from copy import deepcopy
import hashlib, json, re, shutil
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_BREAK

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT.parent/'v2.2/author_review/full_revision_001'
OUT=ROOT/'submission'; OUT.mkdir(exist_ok=True)
E=ROOT/'evidence'
TITLE='Self-Supervised Transfer Learning for Label-Efficient Multigas Electronic Nose Calibration'
ABSTRACT=('Self-supervised transfer reduces mean absolute concentration error relative to training from scratch by 13.9% for NH3, 21.5% for H2, 27.5% for CO, and 25.7% for ethylene in our primary multigas electronic-nose evaluation. '
'We pretrain a shared temporal encoder using masked sensor-response reconstruction, then fine-tune it with a limited set of labeled concentration pairs. '
'The evaluation uses ten seeds, six held-out identifier partitions for NH3/H2, and three rolling chronological folds for CO/ethylene; the NH3/H2 results use a predefined early-response window. '
'In secondary fixed-split experiments, fine-tuning with 10 labeled pairs gives approximately the NH3 and ethylene errors of scratch training with 20 pairs, a lower mean CO error, and a higher H2 error. '
'These pair counts include validation labels. Sensor-identity and cross-sensor-attention ablations show target-dependent effects, with NH3 the weakest transfer case. '
'In a known-gas-identity drift benchmark, transfer improves concentration regression in partitions 8 and 9 but not clearly in partition 10. '
'The proposed model is not the best full-budget baseline on any of the four raw-response targets. '
'The findings support target-dependent label efficiency for same-system calibration; the NH3/H2 identifier partitions do not establish between-day generalization.')
assert len(ABSTRACT.split())<250
changes=[]

def text(p):
    return ''.join(p._p.itertext()) if False else p.text

def replace(p,value,reason):
    old=p.text
    rpr=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for child in list(p._p):
        if child.tag!=qn('w:pPr'):p._p.remove(child)
    r=p.add_run(value)
    if rpr is not None:r._r.insert(0,rpr)
    mark(p)
    changes.append({'reason':reason,'old':old,'new':value})

def mark(p):
    # Paragraph/run highlighting is removed for the clean copy.
    for r in p._p.iter(qn('w:r')):
        pr=r.find(qn('w:rPr'))
        if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
        h=pr.find(qn('w:highlight'))
        if h is None:h=OxmlElement('w:highlight');pr.append(h)
        h.set(qn('w:val'),'yellow')
    for r in p._p.iter(qn('m:r')):
        pr=r.find(qn('w:rPr'))
        if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
        h=OxmlElement('w:highlight');h.set(qn('w:val'),'yellow');pr.append(h)

def add_before(doc,anchor,value,style):
    node=OxmlElement('w:p');anchor.addprevious(node)
    p=Paragraph(node,doc._body);p.style=style
    replace(p,value,'Added revision paragraph')
    return p

def table(doc,anchor,caption,headers,rows,style,width=3.45):
    cap=add_before(doc,anchor,caption,style);cap.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    anchor.addprevious(t._tbl)
    for cell,v in zip(t.rows[0].cells,headers):cell.text=v
    for row in rows:
        for cell,v in zip(t.add_row().cells,row):cell.text=str(v)
    t.autofit=False
    for col in t.columns:col.width=Inches(width/len(headers))
    for i,row in enumerate(t.rows):
        for c in row.cells:
            c.width=Inches(width/len(headers))
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.keep_with_next=i<len(t.rows)-1
                for r in p.runs:r.font.size=Pt(8);r.font.name='Times New Roman';r.bold=i==0
                mark(p)
    return t

def allparagraphs(doc):
    return [Paragraph(n,doc._body) for n in doc.element.body.iter(qn('w:p'))]

def voice(doc):
    # Replace across run boundaries while preserving equations, fields and drawings.
    pairs=[("The authors'",'Our'),("the authors'",'our'),('The authors','We'),('the authors','we'),
           ('avoids need for','avoids the need for'),('their NH3/H2','our NH3/H2')]
    for p in allparagraphs(doc):
        for a,b in pairs:
            nodes=list(p._p.iter(qn('w:t')))
            joined=''.join(n.text or '' for n in nodes)
            while a in joined:
                start=joined.index(a);end=start+len(a);offset=0
                for n in nodes:
                    v=n.text or '';lo,hi=offset,offset+len(v);offset=hi
                    if hi<=start or lo>=end:continue
                    n.text=v[:max(0,start-lo)]+(b if lo<=start<hi else '')+v[max(0,end-lo):]
                mark(p)
                joined=''.join(n.text or '' for n in nodes)

def clear_highlights(doc):
    for h in list(doc.element.iter(qn('w:highlight'))):h.getparent().remove(h)

def save_pair(doc,stem):
    doc.save(OUT/(stem+'_Highlighted.docx'))
    clean=deepcopy(doc);clear_highlights(clean);clean.save(OUT/(stem+'_Clean.docx'))

d=Document(BASE/'IEEE_Full_Revision_Clean.docx');p=list(d.paragraphs)
original_tables=[[ [c.text for c in row.cells] for row in t.rows] for t in d.tables[1:]]
styles={'body':p[120].style,'heading':p[154].style,'caption':p[121].style}
replace(p[0],TITLE,'Title: target-dependent architecture evidence')
abstract=next(q for q in allparagraphs(d) if q.text.startswith('Abstract'))
replace(abstract,'Abstract— '+ABSTRACT,'Result-led abstract below 250 words')
revisions={
17:'Our contribution is a controlled evaluation of masked self-supervised transfer for quantitative same-system recalibration, with limited labels, held-out identifier partitions, chronological tests and direct competing controls.',
62:'Here b indexes examples in a minibatch, t the temporal position, s the physical sensor, and k one of three spans sampled independently for each example. Given span width w_k, its start a_k is sampled uniformly from the integers 0 through T − w_k; the example index on a_k and w_k is suppressed. B_bks denotes sensor participation in span k. M is one for observed entries and zero for hidden entries.',
76:'The NH3/H2 dataset [29] contains 162,404 observations from 20 Nano MEMS sensor channels, grouped into 378 recorded sequences, with concentrations of 0–500 ppm for each gas. Acquisition took place over six days. The six partitions of 63 sequences were constructed from consecutive Combination identifiers and were not designed as day splits. Day-by-day mixture details were not logged, so a reliable exposure-to-day mapping and leave-one-day-out evaluation cannot be reconstructed. Training and test partitions may share acquisition days; these folds do not establish between-day generalization.',
97:'The NH3/H2 array uses 20 MEMS SnO2 channels with Pt/Ag modification, as described in Supplementary “NH3/H2 dataset and acquisition details.” Temperature and relative humidity were neither controlled nor measured during acquisition, so their variation and contribution to response differences cannot be quantified retrospectively. Digital audits found no identical stored signal fingerprints or identifier overlap across 11 split definitions; they do not establish session independence. The 0–500 ppm range does not validate low-concentration ambient monitoring.',
98:'Each NH3/H2 sequence starts at gas delivery and contains 180 s of exposure followed by 700 s of recovery. The configured sampling interval is 2.0 s. The predefined 64-sample early window covers about 128 s nominally (approximately 131 s from aggregate sample counts), or roughly 71–73% of exposure. Middle, late and full-cycle sensitivity inputs also have 64 samples, with scalers refitted on their pretraining windows. Full-cycle resampling includes recovery, rather than only the 180-s gas exposure. Exact index definitions and the approximate timing calculation are supplied in the supplement.',
115:'Table I compares fine-tuning with scratch under matched adaptation groups and holdouts. The four unadjusted paired 95% intervals exclude zero. In a four-comparison Bonferroni sensitivity analysis, the NH3 interval includes zero; the other three remain positive (Supplementary Table S9). NH3 is therefore the weakest primary case. CO/ethylene intervals are indicative because they resample only three chronological folds. Inference is conditional on observed exposures, not independent acquisition campaigns.',
116:'Table I. Primary Multi-Fold MAE and Paired Difference (ppm); Unadjusted 95% Intervals',
121:'Table II. Fixed-Split MAE (ppm), Farthest-Point Calibration-Pair Design, Mean ± SD Over Ten Seeds',
125:'Calibration-design effects depend on both target and budget (Supplementary Table S3). At five pairs, farthest-point design reduces fine-tuned CO MAE relative to random by 31.61 ppm (unadjusted 95% interval 13.34–51.92). At 40 pairs, random design instead has lower MAE: random minus farthest-point differences are −4.46 ppm for NH3 and −8.32 ppm for H2 under scratch training, and −3.37 and −4.89 ppm under fine-tuning; all four exploratory intervals exclude zero. These are known-coordinate calibration designs, not active acquisition of unknown labels. Full-budget designs coincide when all candidate pairs are used.',
130:'The proposed model has the lowest mean MAE on none of the four targets at full budget (Table IV). Generic SSL leads on NH3 (16.13 versus 17.36 ppm), SVR on H2 (31.37 versus 31.60 ppm), Extra Trees on ethylene (1.398 versus 1.627 ppm), and PLSR on CO (25.57 versus 47.45 ppm). Proposed CO MAE is 1.86 times PLSR MAE despite improvement over scratch. A lower-complexity latent linear model may better suit the small calibration set and this response representation; this is a possible explanation, not an identified mechanism. Close mean rankings are not equivalence tests, and overlapping SDs do not establish statistical ties. Tuning grids and all-budget results are supplied in the supplement.',
138:'Window placement materially changes both absolute error and transfer effects (Table VI). Fixed-split early-window H2 improves from 41.20 to 31.60 ppm (23.3%); middle-window transfer reverses the benefit (33.99 to 35.60 ppm, −4.7%); full-cycle resampling yields 38.36 to 35.75 ppm (6.8%). The 21.5% headline instead comes from primary multi-fold evaluation using the same predefined early-window definition. It should not be generalized to other window placements. Middle/late windows sample the recovery part of the recorded cycle; full-cycle resampling spans exposure and recovery. No window is selected using test performance.',
155:'Masked self-supervised pretraining reduces mean error against scratch in this same-system calibration evaluation. Under predefined fine-tuning, primary reductions are 13.9% for NH3, 21.5% for H2, 27.5% for CO and 25.7% for ethylene; NH3/H2 uses the early-response window. Fixed-split pair-count comparisons show approximately comparable NH3/ethylene errors and lower mean CO error with 10 transfer pairs versus 20 scratch pairs, but not an H2 match at those budgets.',
156:'The proposed model is not the best full-budget baseline on any of the four targets; CO MAE is 1.86 times that of PLSR. Sensor identity and attention are not uniformly beneficial, particularly for NH3. NH3 also has negligible five-pair improvement and is the weakest multiplicity-adjusted primary case. Window and calibration-design effects vary with target and budget. Remaining absolute and tail errors limit quantitative usefulness.',
158:'Ten seeds quantify training/design variability, not independent acquisitions. NH3’s four-comparison adjusted interval includes zero; CO/ethylene intervals remain indicative with only three folds. NH3/H2 acquisition spanned six days, but missing day-by-day mixture logs prevent reliable day-held-out evaluation. Temperature and humidity were neither controlled nor measured. Consequently these identifier folds cannot establish between-day or environmental robustness. Cross-device transfer remains unevaluated.',
159:'Label withholding tests the learning premise retrospectively, not acquisition-cost savings. Known-coordinate calibration design excludes candidate-pool quantification costs. Future experiments will measure and record environmental conditions; exposure-linked day/session logs are needed for independent-day evaluation. Prospective unquantified-mixture pretraining, acquisition-cost accounting, and shift-aware uncertainty remain future work.',
}
for i,value in revisions.items():replace(p[i],value,'R3.0 content revision')
# Preserve native editable equations while fixing the symbol collision and guard.
for i in [57,61]:
    for node in p[i]._p.iter(qn('m:t')):
        if node.text=='Z':node.text='B'
    mark(p[i])
den=list(p[65]._p.iter(qn('m:den')))[1]
nodes=list(den.iter(qn('m:t')));nodes[0].text='max(1, ∑';nodes[-1].text='M)';mark(p[65])
# Repair inherited OMML tokens that LibreOffice otherwise parses as commands.
# The mathematical meanings and equation numbering are unchanged.
for node in p[67]._p.iter(qn('m:t')):
    if node.text=='in':
        run=node.getparent();pr=run.find(qn('m:rPr'))
        if pr is None:pr=OxmlElement('m:rPr');run.insert(0,pr)
        for child in list(pr):pr.remove(child)
        pr.append(OxmlElement('m:nor'))
for node in p[72]._p.iter(qn('m:t')):
    if node.text and '\\' in node.text:node.text=node.text.replace('\\','∖')
mark(p[67]);mark(p[72])
add_before(d,p[69]._p,'Both denominators in (14) are guarded against empty sets. In the present raw-response mask, at most 48 of 64 positions are covered by three spans and at most three of 16 or 20 sensors are fully hidden, so visible entries always remain. The guard does not change this protocol’s loss.',styles['body'])
add_before(d,p[109]._p,'Primary intervals in Table I are unadjusted. As a post-review multiplicity sensitivity check, we also report 98.75% marginal percentile intervals for a four-comparison Bonferroni family, using 200,000 crossed seed/fold draws. A second random-number seed checks numerical stability. This adjustment does not remedy dependent training folds or missing acquisition-day labels (Supplementary Table S9).',styles['body'])
add_before(d,p[119]._p,'Transfer reverses in identifier fold 4: NH3 MAE increases from 16.22 to 18.83 ppm and H2 from 25.70 to 26.36 ppm. Aggregate gains therefore do not imply benefits in every partition. The available records do not support attributing this reversal to a particular day or environmental condition.',styles['body'])
table(d,p[124]._p,'Table IIa. Label-Budget Comparison, Farthest-Point Fixed Split; MAE (ppm), Mean ± SD',
      ['Gas','Transfer: 10 pairs','Scratch: 20 pairs'],
      [['NH3','38.92 ± 4.56','36.84 ± 7.89'],['H2','81.88 ± 19.54','70.73 ± 30.45'],['CO','80.99 ± 27.98','89.30 ± 39.24'],['Ethylene','2.418 ± 0.484','2.465 ± 0.745']],styles['caption'])
add_before(d,p[124]._p,'Table IIa expresses label efficiency in labeled two-gas concentration pairs, including fitting and validation. At 10 pairs, transfer approximately matches scratch at 20 for NH3 and ethylene and has lower mean CO error, but higher H2 error. These descriptive comparisons are not equivalence tests or a universal 50% label-saving guarantee. Pairs, exposure records and overlapping windows are distinct units.',styles['body'])
add_before(d,p[124]._p,'At five pairs, NH3 shows essentially no mean benefit (86.39 ± 36.11 to 86.02 ± 34.16 ppm). An error near 86 ppm is 17.2% of its 0–500 ppm range and does not demonstrate useful precise quantitative calibration. Sparse calibration coverage and training variability are possible contributors, not established causes. This is consistent with NH3’s weak primary and component-ablation evidence.',styles['body'])
add_before(d,p[137]._p,'CO ablations are variable: removing attention gives 10.27 ± 12.18 ppm and removing temporal masking 7.93 ± 14.86 ppm. Exploratory paired-seed 95% bootstrap intervals are [2.68, 16.94] and [−0.33, 17.00] ppm, respectively (Supplementary Table S10). Temporal masking is compatible with no effect; attention has a positive conditional mean contrast, but its large variability limits a robust component-benefit claim. Neither SD overlap nor these unadjusted intervals establishes a general physical mechanism.',styles['body'])
# Explicitly label the full-cycle row; retain every value.
window_table=next(t for t in d.tables if t.cell(0,0).text=='Window')
for row in window_table.rows:
    if row.cells[0].text=='Full':replace(row.cells[0].paragraphs[0],'Full-cycle','Window nomenclature')
timing=json.loads((E/'compute_measurement.json').read_text())
times={r['dataset']:r['seconds'] for r in timing['runs']}
compute=f'Representative seed-42 pretraining reruns on one {timing["gpu"]} took {times["nh3_h2"]:.1f} s for NH3/H2 (1,260 windows) and {times["co_ethylene"]:.1f} s for CO/ethylene (735 windows), each for 60 epochs with batch size 64. Timings include model initialization and optimization but exclude data loading, scaling and checkpoint writes. They are per-model measurements, not reconstructed total campaign costs; checkpoints are reused across adaptation budgets. Software and timing details are supplied in the supplement.'
add_before(d,p[101]._p,compute,styles['body'])
add_before(d,p[160]._p,'Data and Code Availability',styles['heading'])
add_before(d,p[160]._p,'The NH3/H2 dataset is identified by Harvard Dataverse DOI 10.7910/DVN/LWUSKO [29]; the public CO/ethylene and drift datasets are cited in [15] and [16]. Study code is linked at https://github.com/ngdvan/Transfer-learning-Enose. The accompanying revision package contains the analysis scripts, split-audit records, unrounded result exports and representative compute measurements used here. Exposure-to-day mixture logs and environmental measurements were not recorded and are not available.',styles['body'])
p[200]._p.getparent().remove(p[200]._p)
# Keep the first biography photo with its text in the next column.
bio_break=add_before(d,p[191]._p,'',styles['body'])
bio_break.add_run().add_break(WD_BREAK.COLUMN)
voice(d)
save_pair(d,'IEEE_Revision_3.0')
(OUT/'ABSTRACT.md').write_text('# '+TITLE+'\n\n'+ABSTRACT+'\n\nWord count (whitespace-delimited): '+str(len(ABSTRACT.split()))+'\n',encoding='utf-8')

# Supplement: preserve the existing evidence tables and hardware images.
s=Document(BASE/'SUPPLEMENTARY_EVIDENCE.docx');sp=list(s.paragraphs)
replace(sp[27], 'We programmed single-gas and binary NH3/H2 conditions spanning 0–500 ppm for each gas, with nonzero settings of 25, 50, 100, 200, 300, 400 and 500 ppm. Each sequence starts at gas delivery and records the complete 180-s exposure plus 700-s recovery cycle (880 s total). The acquisition interface was configured for a 2.0-s sampling interval. Acquisition took place over six days, but the six Combination-ID partitions were not intended as day splits. Day-by-day mixture details were not logged, preventing a reliable exposure-to-day mapping and leave-one-day-out evaluation. Temperature and relative humidity were neither controlled nor measured. These facts are clarified by the acquisition team; no timestamp-based cadence validation or retrospective environmental measurements are claimed.', 'Author-confirmed acquisition protocol')
voice(s)
s.add_heading('Revision 3.0: additional analysis and protocol clarification',level=1)
s.add_heading('Acquisition independence and environmental limitations',level=2)
s.add_paragraph('Identifier separation and stored-signal fingerprint checks establish digital separation only. Acquisition days may be shared across evaluation roles; without exposure-to-day records, overlap cannot be quantified. The existing folds support held-out identifier evaluation, not verified day independence. Environmental variation cannot be quantified retrospectively. Future experiments will measure and record temperature and humidity; exposure-linked day/session logs are needed to enable day-held-out validation.')
s.add_heading('Physical timing and exact window definitions',level=2)
s.add_paragraph('The configured interval is 2.0 s. Aggregate counts give 162,404/378 = 429.64 observations per 880-s sequence, corresponding approximately to 2.048 s per observation. This is consistent with the configured interval, but is not a timestamp measurement. A 64-sample window represents about 128 s nominally or 131 s by this aggregate calculation (71–73% of the 180-s exposure); the first-to-last timestamp span at nominal cadence is 63 intervals, or 126 s. The first sample marks gas-delivery onset.')
s.add_paragraph('For an L-sample sequence, with zero-based indices: early windows start at offsets 0–4; middle windows start at floor((L−68)/2) plus offsets 0–4; late windows start at L−68 plus offsets 0–4. Each has 64 consecutive samples. Tests use only offset zero. Full-cycle windows linearly interpolate the complete interval [offset, L−1] to 64 points, using offsets 0–4 for pretraining/fitting and zero for testing. Full-cycle inputs include gas exposure and recovery. For approximately 430-sample cycles, middle and late windows predominantly probe recovery. Scalers are refitted using the corresponding pretraining windows; the primary early window remains predefined.')
s.add_heading('Multiplicity sensitivity',level=2)
s.add_paragraph('Table S9 retains equal fold weighting and independently resamples paired seed and fold axes. The 98.75% marginal percentile intervals implement a Bonferroni sensitivity for four comparisons using 200,000 draws (RNG seed 42). A second 200,000-draw run with RNG seed 2026 gives the same zero-inclusion decisions; endpoint shifts are at most 0.15 ppm. These are conditional descriptive bootstrap intervals. Bonferroni adjustment does not repair the limited number of chronological folds, overlapping training data, or unavailable acquisition-day labels. The original unadjusted Table I remains unchanged.')
adj=pd.read_csv(E/'multiplicity_sensitivity.csv');adj=adj[adj.rng_seed==42].set_index('gas')
s.add_paragraph('Table S9. Four-comparison multiplicity sensitivity; scratch minus fine-tuned MAE (ppm).')
t=s.add_table(rows=1,cols=4);t.style='Table Grid'
for c,v in zip(t.rows[0].cells,['Gas','Difference','98.75% interval','Test folds']):c.text=v
for gas in ['NH3','H2','CO','Ethylene']:
    r=adj.loc[gas];digits=3 if gas=='Ethylene' else 2
    for c,v in zip(t.add_row().cells,[gas,f'{r.difference:.{digits}f}',f'[{r.adjusted_low:.{digits}f}, {r.adjusted_high:.{digits}f}]',str(int(r.folds))]):c.text=v
s.add_paragraph('NH3 includes zero after adjustment; H2, CO and ethylene remain positive. CO/ethylene intervals are indicative with only three test folds. NH3 is also the weakest case at five pairs and for sensor-ID/attention ablations.')
s.add_heading('CO ablation uncertainty',level=2)
s.add_paragraph('Table S10. Full-budget CO ablation-minus-full-model MAE: exploratory paired-seed percentile bootstrap, 200,000 draws, RNG seed 42, unadjusted 95% intervals. Seeds remain paired; test exposures are not resampled as independent acquisitions.')
t=s.add_table(rows=1,cols=3);t.style='Table Grid'
for c,v in zip(t.rows[0].cells,['Removed','Mean ± SD (ppm)','95% interval (ppm)']):c.text=v
for _,r in pd.read_csv(E/'co_ablation_sensitivity.csv').iterrows():
    for c,v in zip(t.add_row().cells,[{'no_attention':'Attention','no_temporal_mask':'Temporal mask'}[r.variant],f'{r["mean"]:.2f} ± {r.sd:.2f}',f'[{r.ci_low:.2f}, {r.ci_high:.2f}]']):c.text=v
s.add_paragraph('The temporal-mask contrast includes zero. The attention contrast is positive in this conditional analysis, so it should not be called a statistical null solely because its mean is smaller than one SD. Both have substantial variability and neither supports a universal component benefit.')
s.add_heading('Representative pretraining compute',level=2)
s.add_paragraph(compute)
s.add_paragraph(f'Timing environment: Python {timing["python"]}, PyTorch {timing["torch"]}, CUDA {timing["cuda"]}. One GPU was used. The reported reruns use the original pretraining implementation and fixed-split preprocessing, without adaptation training or reuse of a cached pretrained state. The combined measured time for these two representative runs is {sum(times.values()):.1f} s; it is not the total ten-seed/multi-fold experiment cost. Descriptor-benchmark pretraining uses 40 epochs (batch size 256); its wall time was not measured in this raw-response timing check. Hardware/software and source hashes are saved in evidence/compute_measurement.json.')
s.add_heading('Loss denominator guard',level=2)
s.add_paragraph('Equation (14) guards both denominators. The V3.0 raw-response implementation adds clamp_min(1) to the visible denominator. With at most three spans of length 16 in a 64-sample window and at most three entirely hidden channels, at least 16 temporal positions remain visible in each nonhidden channel. The visible denominator is therefore always positive for the reported 16-/20-channel protocols. This defensive code change leaves their loss unchanged; historical predictions and results are retained.')
clear_highlights(s);s.save(OUT/'SUPPLEMENTARY_EVIDENCE.docx')

# Preserve versioned experiment code and original outputs without touching V2.2.
shutil.copytree(BASE/'evidence/exports',E/'exports',dirs_exist_ok=True)
shutil.copytree(ROOT.parent/'v2.2/code',ROOT/'code',dirs_exist_ok=True,ignore=shutil.ignore_patterns('__pycache__'))
code=ROOT/'code/same_device_session_transfer.py'
src=code.read_text(encoding='utf-8')
src=src.replace('/ observed.sum()','/ observed.sum().clamp_min(1)')
src=src.replace('Sessions 1-4:','Identifier partitions 1-4:').replace('Session 5:','Identifier partition 5:').replace('Session 6: untouched test session.','Identifier partition 6: held-out test exposures (not a day split).')
code.write_text(src,encoding='utf-8')
for folder in ['leakage_audit','hardware']:
    shutil.copytree(BASE/'evidence'/folder,E/folder,dirs_exist_ok=True)
# Historical acquisition notes are source records, not current author-confirmed claims.
(E/'SOURCE_NOTES.md').write_text('The exports and hardware/leakage notes copied from V2.2 are historical source records. Author-confirmed acquisition corrections are in the V3.0 manuscript/supplement and action plan. The variable `sessions` in inherited code denotes identifier partitions, not logged acquisition days.\n',encoding='utf-8')
(E/'document_changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
manifest={str(f.relative_to(ROOT.parent)):hashlib.sha256(f.read_bytes()).hexdigest() for f in [BASE/'IEEE_Full_Revision_Clean.docx',BASE/'SUPPLEMENTARY_EVIDENCE.docx',*(ROOT/'Base v2.2').glob('*.pdf')]}
(E/'source_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print('Built manuscript pair and supplement. Abstract words:',len(ABSTRACT.split()))
